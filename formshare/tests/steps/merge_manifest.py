"""A wide form's new version takes submissions after a merge, and an answer
with an emoji can be changed in a repository built or merged since
createaudittriggers writes audit_log utf8mb4.

A merge keeps a split table's columns where the repository has them, and adds
a new one at the end of the chain, whatever the new version's own split says
(rstools.md 36.5). The version's own manifest follows its own split, and the
version's submissions used to be loaded through it: an answer it placed in a
table that does not hold it refused the whole submission (Unknown column).
mergeversions now rewrites the manifest to follow the repository (-m, -M),
the merge makes that the form's, and a device is served it with the
create.xml and insert.xml the repository was built from.

The step builds a form of two hundred selects, which splits, and merges into
it a version with one more question ahead of the split: the version's own
manifest puts it in maintable, the merge at the end of the chain. Since
stage 8, _assigned_to moves a wide version's split the same way when its
repository was built before it.

createaudittriggers writes audit_log utf8mb4 since rstools.md 37 and converts
an older one whenever it runs, which a build and every merge now do. A merge
used to run it only when a table gained a field: one that only added a table
left that table unaudited, and one that changed no table never ran it. The
emoji checks run when the installed tool knows it (CI's image may predate
it); the merges are checked either way. Existing repositories are not
regenerated (owner, 2026-10-17): they take emoji from their next merge.
"""

import json
import os
import shutil
import uuid
from subprocess import PIPE, Popen

import openpyxl
from lxml import etree
from sqlalchemy import create_engine
from sqlalchemy.pool import NullPool

import formshare.processes.odk.api as odk_api
from .in_process import add_task, build_in_process, form_row

FORM = "wide_merge"
VERSION = "wide_merge_v2"
REPEAT = "wide_merge_v3"
UNCHANGED = "wide_merge_v4"
SELECTS = 200
EMOJI = "Lote con rebrote, sin producción, 💪😊"
SURVEY_COLUMNS = [
    "formshare_sensitive",
    "formshare_encrypted",
    "formshare_ontological_term",
    "formshare_unique",
]


def _query(config, sql, params=None):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        if params is None:
            return engine.execute(sql).fetchall()
        return engine.execute(sql, params).fetchall()
    finally:
        engine.dispose()


def _execute(config, sql):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        engine.execute(sql)
    finally:
        engine.dispose()


def tool_converts_audit_log(config, work):
    """True when the installed createaudittriggers converts an audit_log
    older than rstools.md 37 to utf8mb4, False when it leaves it as it is.
    Asks it about a scratch schema whose audit_log is utf8mb3, naming that
    table alone, as a merge that changes no table does."""
    schema = "fs_probe_" + uuid.uuid4().hex[:12]
    out_dir = os.path.join(work, "audit_probe")
    os.makedirs(out_dir, exist_ok=True)
    _execute(config, "CREATE SCHEMA `{}`".format(schema))
    try:
        _execute(
            config,
            "CREATE TABLE `{}`.audit_log (audit_id VARCHAR(64) PRIMARY KEY, "
            "audit_newvalue TEXT) CHARSET=utf8mb3".format(schema),
        )
        process = Popen(
            [
                os.path.join(
                    config["odktools.path"],
                    "utilities",
                    "createAuditTriggers",
                    "createaudittriggers",
                ),
                "-H " + config["mysql.host"],
                "-P " + str(config["mysql.port"]),
                "-u " + config["mysql.user"],
                "-p " + config["mysql.password"],
                "-s " + schema,
                "-o " + out_dir,
                "-t audit_log",
            ],
            stdout=PIPE,
            stderr=PIPE,
        )
        process.communicate()
        sql_file = os.path.join(out_dir, "mysql_create_audit.sql")
        if process.returncode != 0 or not os.path.exists(sql_file):
            return False
        with open(sql_file) as a_file:
            return "CONVERT TO CHARACTER SET utf8mb4" in a_file.read()
    finally:
        _execute(config, "DROP SCHEMA IF EXISTS `{}`".format(schema))
        shutil.rmtree(out_dir, ignore_errors=True)


def _workbook(path, form_id, version, early_question, repeat=False):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "i_d", "ID"])
    if early_question:
        # Ahead of the split: the version's own build keeps it in maintable,
        # and a merge adds it at the end of the repository's chain
        survey.append(["text", "early_note", "Early note"])
    survey.append(["text", "remarks", "Remarks"])
    for i in range(1, SELECTS + 1):
        survey.append(["select_one l%03d" % i, "q%03d" % i, "Question %d" % i])
    survey.append(["integer", "late_count", "Late count"])
    if repeat:
        # A table of its own and nothing else new
        survey.append(["begin_repeat", "visits", "Visits"])
        survey.append(["text", "visit_note", "Visit note"])
        survey.append(["end_repeat", "", ""])
    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    for i in range(1, SELECTS + 1):
        for j in (1, 2):
            choices.append(["l%03d" % i, "a%d_%d" % (i, j), "Option %d of %d" % (j, i)])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Wide form, merged", form_id, version])
    wb.save(path)


def _submission(path, form_id, key, early_note=None, visit_note=None):
    data = {
        "_xform_id_string": form_id,
        "i_d": key,
        "remarks": "plain",
        "late_count": "4",
        "meta/instanceID": "uuid:{}".format(uuid.uuid4()),
    }
    for i in range(1, SELECTS + 1):
        data["q%03d" % i] = "a%d_1" % i
    if early_note is not None:
        data["early_note"] = early_note
    if visit_note is not None:
        data["visits"] = [{"visits/visit_note": visit_note}]
    with open(path, "w", encoding="utf-8") as a_file:
        json.dump(data, a_file)


def _placements(manifest_file):
    """Where a manifest puts each answer: {column: table}. The keys are left
    out: every table of a chain carries a copy of them."""
    placed = {}
    for table in etree.parse(manifest_file).iter("table"):
        for field in table.findall("field"):
            if field.get("key") == "true" or field.get("xmlcode") in (None, "NONE"):
                continue
            placed[field.get("mysqlcode")] = table.get("mysqlcode")
    return placed


def _holds(config, schema, table, column):
    return (
        _query(
            config,
            "SELECT COUNT(*) FROM information_schema.COLUMNS WHERE TABLE_SCHEMA = %s "
            "AND TABLE_NAME = %s AND COLUMN_NAME = %s",
            (schema, table, column),
        )[0][0]
        == 1
    )


def _add_assistant(test_object, form_id):
    res = test_object.testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(
            test_object.randonLogin, test_object.project, form_id
        ),
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )
    assert "FS_error" not in res.headers


def _as_the_assistant(test_object):
    return dict(FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin)


def _push(test_object, path):
    res = test_object.testapp.post(
        "/user/{}/project/{}/push_json".format(
            test_object.randonLogin, test_object.project
        ),
        status=201,
        upload_files=[("filetoupload", path)],
        extra_environ=_as_the_assistant(test_object),
    )
    assert "FS_error" not in res.headers


def _upload_version(test_object, work, form_id, version, parent, repeat=False):
    xlsx = os.path.join(work, form_id + ".xlsx")
    _workbook(xlsx, form_id, version, early_question=True, repeat=repeat)
    res = test_object.testapp.post(
        "/user/{}/project/{}/form/{}/merge".format(
            test_object.randonLogin, test_object.project, parent
        ),
        {
            "for_merging": "",
            "parent_project": test_object.projectID,
            "parent_form": parent,
        },
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    _add_assistant(test_object, form_id)


def _merge(test_object, version, parent):
    """The merge check, as the version's page opens, and then the merge in
    this process, as the worker runs it. Returns the version's row."""
    from formshare.products.merge.celery_task import internal_merge_into_repository

    login = test_object.randonLogin
    project = test_object.project
    project_id = test_object.projectID
    config = test_object.server_config
    res = test_object.testapp.get(
        "/user/{}/project/{}/form/{}".format(login, project, version), status=200
    )
    assert b"Merge check pending" not in res.body
    test_object.testapp.get(
        "/user/{}/project/{}/form/{}/merge/into/{}".format(
            login, project, version, parent
        ),
        status=200,
    )
    a_form = form_row(config, project_id, version)
    b_form = form_row(config, project_id, parent)
    internal_merge_into_repository(
        config,
        login,
        project_id,
        project,
        version,
        a_form["form_directory"],
        b_form["form_directory"],
        b_form["form_createxmlfile"],
        b_form["form_schema"],
        "",
        b_form["form_hexcolor"],
        "en",
        True,
        SURVEY_COLUMNS,
        add_task(config, project_id, version, "merge_form", login),
    )
    a_form = form_row(config, project_id, version)
    assert a_form["form_schema"] == b_form["form_schema"]
    assert a_form["form_blocked"] == 0
    return a_form


def _edit_remarks(test_object, form_id, rowuuid, value):
    """The answer called remarks edited in the cleaning grid by the assistant.
    Returns what the grid is told: {} when the value is saved."""
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    res = testapp.post(
        "/user/{}/project/{}/assistantaccess/login".format(login, project),
        {"login": test_object.assistantLogin, "passwd": "123"},
        status=302,
    )
    assert "FS_error" not in res.headers
    try:
        res = testapp.post(
            "/user/{}/project/{}/assistantaccess/form/{}/maintable/action".format(
                login, project, form_id
            ),
            {"remarks": value, "oper": "edit", "id": rowuuid},
            status=200,
        )
        return res.json
    finally:
        testapp.post(
            "/user/{}/project/{}/assistantaccess/logout".format(login, project),
            status=302,
        )


def _remarks(config, schema, rowuuid):
    return _query(
        config,
        "SELECT remarks FROM `{}`.maintable WHERE rowuuid = %s".format(schema),
        (rowuuid,),
    )[0][0]


def _audit_log_collation(config, schema):
    return _query(
        config,
        "SELECT TABLE_COLLATION FROM information_schema.TABLES "
        "WHERE TABLE_SCHEMA = %s AND TABLE_NAME = 'audit_log'",
        (schema,),
    )[0][0]


def _make_audit_log_older(config, schema):
    """audit_log as createaudittriggers wrote it before rstools.md 37."""
    _execute(
        config,
        "ALTER TABLE `{}`.audit_log CONVERT TO CHARACTER SET utf8mb3 "
        "COLLATE utf8mb3_general_ci".format(schema),
    )


def _make_audit_log_older_again(config, schema):
    """utf8mb3 cannot hold what the emoji edits logged, so their rows go
    first."""
    _execute(
        config,
        "DELETE FROM `{}`.audit_log WHERE audit_column = 'remarks'".format(schema),
    )
    _make_audit_log_older(config, schema)


def _triggers(config, schema):
    return _query(
        config,
        "SELECT COUNT(*) FROM information_schema.TRIGGERS WHERE TRIGGER_SCHEMA = %s",
        (schema,),
    )[0][0]


def _check_an_emoji_is_saved(test_object, form_id, schema, rowuuid, value):
    config = test_object.server_config
    assert _audit_log_collation(config, schema) == "utf8mb4_unicode_ci"
    assert _edit_remarks(test_object, form_id, rowuuid, value) == {}
    assert _remarks(config, schema, rowuuid) == value
    assert (value,) in _query(
        config,
        "SELECT audit_newvalue FROM `{}`.audit_log WHERE audit_action = 'UPDATE' "
        "AND audit_column = 'remarks'".format(schema),
    )


def t_e_s_t_merge_manifest(test_object):
    login = test_object.randonLogin
    project = test_object.project
    project_id = test_object.projectID
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "merge_manifest")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)
    emoji = tool_converts_audit_log(config, work)

    # --- The wide form, its repository, a submission -----------------------------
    xlsx = os.path.join(work, FORM + ".xlsx")
    _workbook(xlsx, FORM, "1", early_question=False)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "i_d"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    _add_assistant(test_object, FORM)
    build_in_process(test_object, project_id, project, FORM, "i_d", [])
    schema = form_row(config, project_id, FORM)["form_schema"]
    assert schema, "the wide form's repository did not build"
    assert _query(
        config,
        "SELECT COUNT(*) FROM information_schema.TABLES WHERE TABLE_SCHEMA = %s "
        "AND TABLE_NAME = 'maintable_ext1'",
        (schema,),
    ) == [(1,)], "the wide form did not split"
    path = os.path.join(work, "W1.json")
    _submission(path, FORM, "W1")
    _push(test_object, path)
    first = _query(
        config, "SELECT rowuuid FROM `{}`.maintable WHERE i_d = 'W1'".format(schema)
    )[0][0]

    # --- audit_log, in a new repository and in an older one ----------------------
    if emoji:
        assert _audit_log_collation(config, schema) == "utf8mb4_unicode_ci"
        # An older repository's cannot take an answer with an emoji: the edit
        # is refused whole
        _make_audit_log_older(config, schema)
        result = _edit_remarks(test_object, FORM, first, EMOJI)
        assert [a_error["field"] for a_error in result.get("errors", [])] == [
            "remarks"
        ], result
        assert _remarks(config, schema, first) == "plain"

    # --- A version with a question ahead of the split, merged --------------------
    _upload_version(test_object, work, VERSION, "2", FORM)
    version = _merge(test_object, VERSION, FORM)
    form_dir = os.path.join(
        config["repository.path"], "odk", "forms", version["form_directory"]
    )
    own = _placements(os.path.join(form_dir, "repository", "manifest_own.xml"))
    merged = _placements(os.path.join(form_dir, "repository", "manifest.xml"))
    with open(os.path.join(form_dir, "merging_files", "manifest.xml"), "rb") as a_file:
        merged_manifest = a_file.read()
    with open(os.path.join(form_dir, "repository", "manifest.xml"), "rb") as a_file:
        assert a_file.read() == merged_manifest
    # The version's own manifest would send the answer where the repository
    # does not keep it; the merged one sends it where it does
    assert own["early_note"] == "maintable"
    assert not _holds(config, schema, "maintable", "early_note")
    assert merged["early_note"] != "maintable"
    assert _holds(config, schema, merged["early_note"], "early_note")

    path = os.path.join(work, "W2.json")
    _submission(path, VERSION, "W2", early_note="Noted early")
    _push(test_object, path)
    assert _query(
        config,
        "SELECT early_note FROM `{}`.`{}` WHERE i_d = 'W2'".format(
            schema, merged["early_note"]
        ),
    ) == [("Noted early",)]

    # --- What a device is served -------------------------------------------------
    # The repository files reach devices in projects that publish a list
    with open(version["form_createxmlfile"], "rb") as a_file:
        merged_create = a_file.read()
    with open(os.path.join(form_dir, "merging_files", "insert.xml"), "rb") as a_file:
        merged_insert = a_file.read()
    uses_registry = odk_api.project_uses_registry
    odk_api.project_uses_registry = lambda request, a_project: True
    try:

        def served(form_id, file_name):
            return testapp.get(
                "/user/{}/project/{}/{}/manifest/repository/{}".format(
                    login, project, form_id, file_name
                ),
                status=200,
                extra_environ=_as_the_assistant(test_object),
            ).body

        # Every version of the repository is served the create.xml and the
        # insert.xml it was built from; the version merged in, the manifest
        # that follows them
        for form_id in (VERSION, FORM):
            assert served(form_id, "create.xml") == merged_create
            assert served(form_id, "insert.xml") == merged_insert
        assert served(VERSION, "manifest.xml") == merged_manifest
    finally:
        odk_api.project_uses_registry = uses_registry

    # The merge converted the older audit_log
    if emoji:
        _check_an_emoji_is_saved(test_object, VERSION, schema, first, EMOJI)

    # --- A version that adds a table and nothing else, merged ---------------------
    # The audit triggers used to be written only when a table gained a field:
    # a table a merge added and nothing else was never audited, and the tool
    # never ran to convert audit_log either
    if emoji:
        _make_audit_log_older_again(config, schema)
    _upload_version(test_object, work, REPEAT, "3", VERSION, repeat=True)
    _merge(test_object, REPEAT, VERSION)
    assert _query(
        config,
        "SELECT COUNT(*) FROM information_schema.TRIGGERS WHERE TRIGGER_SCHEMA = %s "
        "AND EVENT_OBJECT_TABLE = 'visits' AND TRIGGER_NAME LIKE 'audit\\_%%'",
        (schema,),
    )[0][0], "the table the merge added has no audit trigger"
    path = os.path.join(work, "W3.json")
    _submission(path, REPEAT, "W3", early_note="Noted again", visit_note="First")
    _push(test_object, path)
    assert _query(
        config,
        "SELECT visit_note FROM `{}`.visits WHERE i_d = 'W3'".format(schema),
    ) == [("First",)]
    if emoji:
        # Its old value holds an emoji too
        _check_an_emoji_is_saved(test_object, REPEAT, schema, first, "again " + EMOJI)

    # --- A version that changes no table, merged ---------------------------------
    if emoji:
        _make_audit_log_older_again(config, schema)
    triggers = _triggers(config, schema)
    _upload_version(test_object, work, UNCHANGED, "4", REPEAT, repeat=True)
    _merge(test_object, UNCHANGED, REPEAT)
    assert _triggers(config, schema) == triggers, "the merge added audit triggers"
    path = os.path.join(work, "W4.json")
    _submission(path, UNCHANGED, "W4", early_note="Noted once more", visit_note="Two")
    _push(test_object, path)
    assert _query(
        config,
        "SELECT early_note FROM `{}`.`{}` WHERE i_d = 'W4'".format(
            schema, merged["early_note"]
        ),
    ) == [("Noted once more",)]
    if emoji:
        _check_an_emoji_is_saved(
            test_object, UNCHANGED, schema, first, "unchanged " + EMOJI
        )
