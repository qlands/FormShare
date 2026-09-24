"""Actions over a form split across continuation tables, with generated
columns and a GeoJSON lookup (rstools.md 13.4), acting on its own rows.

A form too wide for one MySQL row is split by JXFormToMysql into a chain,
maintable then maintable_ext1, each continuation row naming the row it
continues in link_rowuuid. The server's input builder used to take a
continuation for a repeat and order it by a _rowid it does not have (MySQL
1054), and to hand a module the geometry MySQL derives -- a trace's
``<name>_geom``, a GeoJSON lookup's ``geometry`` -- as bytes. RSTools'
MirrorInput reads the device's mirror the right way; this checks the server
reads MySQL the same way:

- a continuation's columns are the submission's (``s.late_count``);
- a multi-select whose junction rows hang off the continuation row is the
  submission's too, its codes in their order;
- a repeat is still a repeat, with its own multi-select;
- nothing MySQL generates reaches the module, in ``s`` or in a lookup, and
  a lookup row has no ``_rowid``;
- a form without a case link acts on its own rows: the case is the
  submission itself, as the device does with CaseLink.own().
"""

import json
import os
import shutil
import time
import uuid

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.pool import NullPool

from .sql import get_form_details

FORM = "wide_actions"
SELECTS = 200


def _workbook(path):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "i_d", "ID"])
    for i in range(1, SELECTS + 1):
        survey.append(["select_one l%03d" % i, "q%03d" % i, "Question %d" % i])
    # After two hundred selects these land in the continuation.
    survey.append(["select_multiple crop_list", "late_crops", "Late crops"])
    survey.append(["integer", "late_count", "Late count"])
    survey.append(["geotrace", "route", "Route"])
    survey.append(["select_one_from_file museums.geojson", "museum", "Museum"])
    survey.append(["begin_repeat", "plots", "Plots"])
    survey.append(["text", "plot_name", "Plot name"])
    survey.append(["decimal", "plot_area", "Plot area"])
    survey.append(["select_multiple use_list", "plot_uses", "Uses"])
    survey.append(["end_repeat", "", ""])
    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    for i in range(1, SELECTS + 1):
        for j in (1, 2):
            choices.append(["l%03d" % i, "a%d_%d" % (i, j), "Option %d of %d" % (j, i)])
    for code in ("c1", "c2", "c3"):
        choices.append(["crop_list", code, "Crop " + code])
    for code in ("u1", "u2", "u3"):
        choices.append(["use_list", code, "Use " + code])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Wide form with actions", FORM, "1"])
    wb.save(path)


def _submission(path, key):
    answers = "".join("<q%03d>a%d_1</q%03d>" % (i, i, i) for i in range(1, SELECTS + 1))
    xml = (
        "<?xml version='1.0'?>"
        '<data id="{form}" version="1">'
        "<i_d>{key}</i_d>{answers}"
        "<late_crops>c3 c1</late_crops><late_count>4</late_count>"
        "<route>0.1 0.2 0 0;0.3 0.4 0 0</route>"
        "<museum>fs87b</museum>"
        "<plots><plot_name>North</plot_name><plot_area>1.5</plot_area>"
        "<plot_uses>u2 u1</plot_uses></plots>"
        "<plots><plot_name>South</plot_name><plot_area>2</plot_area>"
        "<plot_uses>u3</plot_uses></plots>"
        "<meta><instanceID>uuid:{uuid}</instanceID></meta></data>"
    ).format(form=FORM, key=key, answers=answers, uuid=uuid.uuid4())
    with open(path, "w", encoding="utf-8") as a_file:
        a_file.write(xml)


def _query(config, sql):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        return engine.execute(sql).fetchall()
    finally:
        engine.dispose()


def t_e_s_t_actions_split(test_object):
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "actions_split")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)

    # --- The form, split, with its lookup file, built -----------------------
    xlsx = os.path.join(work, FORM + ".xlsx")
    _workbook(xlsx)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "i_d"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    museums = os.path.join(
        test_object.path, "resources", "forms", "GeoJSON", "museums.geojson"
    )
    res = testapp.post(
        "/user/{}/project/{}/form/{}/upload".format(login, project, FORM),
        status=302,
        upload_files=[("filetoupload", museums)],
    )
    assert "FS_error" not in res.headers
    testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(login, project, FORM),
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )
    testapp.post(
        "/user/{}/project/{}/form/{}/repository/create".format(login, project, FORM),
        {"form_pkey": "i_d", "start_stage1": ""},
        status=302,
    )
    schema = None
    for _ in range(100):
        schema = get_form_details(config, test_object.projectID, FORM)["form_schema"]
        if (
            schema
            and _query(
                config,
                "SELECT COUNT(*) FROM information_schema.TABLES WHERE TABLE_SCHEMA = '{}' "
                "AND TABLE_NAME = 'maintable'".format(schema),
            )[0][0]
        ):
            break
        time.sleep(3)
    assert schema, "the wide form's repository did not build"
    tables = {
        r[0]
        for r in _query(
            config,
            "SELECT TABLE_NAME FROM information_schema.TABLES WHERE TABLE_SCHEMA = '{}'".format(
                schema
            ),
        )
    }
    assert "maintable_ext1" in tables, "two hundred selects should split maintable"
    generated = {
        (r[0], r[1])
        for r in _query(
            config,
            "SELECT TABLE_NAME, COLUMN_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = '{}' AND EXTRA LIKE '%GENERATED%'".format(schema),
        )
    }
    lookup = [t for t in tables if t.startswith("lkpmuseum")]
    assert lookup and (lookup[0], "geometry") in generated, generated
    lookup_list = lookup[0][3:]
    assert any(c.endswith("_geom") for _, c in generated), generated

    # --- One submission --------------------------------------------------------
    xml = os.path.join(work, "sub1.xml")
    _submission(xml, "W1")
    res = testapp.post(
        "/user/{}/project/{}/push".format(login, project),
        status=201,
        upload_files=[("filetoupload", xml)],
        extra_environ=dict(
            FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin
        ),
    )
    assert "FS_error" not in res.headers
    main1 = _query(
        config, "SELECT rowuuid FROM `{}`.maintable WHERE i_d = 'W1'".format(schema)
    )[0][0]
    ext_row = _query(
        config,
        "SELECT rowuuid, late_count FROM `{}`.maintable_ext1 WHERE link_rowuuid = '{}'".format(
            schema, main1
        ),
    )
    assert ext_row and ext_row[0][1] == 4, "late_count should be in the continuation"

    # --- The dry run reads the split form the way the device does -------------
    actions = "/user/{}/project/{}/form/{}/actions".format(login, project, FORM)
    reader = (
        "export default function run(s, api) {\n"
        '  api.log("late " + s.late_count + " " + typeof s.late_count);\n'
        '  api.log("crops " + s.selections("late_crops").join(","));\n'
        '  api.log("q001 " + s.q001 + " q200 " + s.q200);\n'
        "  var names = Object.keys(s).filter(function (k) { return /_geom$/.test(k); });\n"
        '  api.log("generated " + names.length + " route " + (s.route !== null));\n'
        '  s.repeat("plots").forEach(function (p) {\n'
        '    api.log("plot " + p.plot_name + " " + p.plot_area + " " + p.selections("plot_uses").join(","));\n'
        "  });\n"
        '  var m = api.lookup("' + lookup_list + '", s.museum);\n'
        '  api.log("museum " + Object.keys(m).sort().join(","));\n'
        '  api.log("case " + (api.case.rowuuid === s.rowuuid) + " " + api.case.table);\n'
        "}\n"
    )
    res = testapp.post_json(
        actions + "/test", {"rowuuid": main1, "module": reader}, status=200
    )
    assert res.json["ok"], res.json
    log_lines = res.json["log"]
    assert log_lines[0] == "late 4 number", log_lines
    assert log_lines[1] == "crops c3,c1", log_lines
    assert log_lines[2] == "q001 a1_1 q200 a200_1", log_lines
    assert log_lines[3] == "generated 0 route true", log_lines
    assert log_lines[4] == "plot North 1.5 u2,u1", log_lines
    assert log_lines[5] == "plot South 2 u3", log_lines
    museum_columns = log_lines[6][len("museum ") :].split(",")
    assert "geometry" not in museum_columns and "geometry_json" in museum_columns
    assert not [c for c in museum_columns if c.endswith("_rowid") or c == "rowuuid"]
    assert log_lines[7] == "case true maintable", log_lines

    # --- A real run on the form's own row, recorded -----------------------------
    res = testapp.post(actions, {"set_mode": "1", "action_mode": "expert"}, status=302)
    assert "FS_error" not in res.headers
    res = testapp.post(
        actions,
        {
            "save_module": "1",
            "action_module": "export default function run(s, api) {\n"
            "  if (s.late_count > 3 && s.selected('late_crops', 'c3')) { api.case.deactivate(); }\n"
            "}\n",
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    xml = os.path.join(work, "sub2.xml")
    _submission(xml, "W2")
    res = testapp.post(
        "/user/{}/project/{}/push".format(login, project),
        status=201,
        upload_files=[("filetoupload", xml)],
        extra_environ=dict(
            FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin
        ),
    )
    assert "FS_error" not in res.headers
    main2, active = _query(
        config,
        "SELECT rowuuid, _active FROM `{}`.maintable WHERE i_d = 'W2'".format(schema),
    )[0]
    assert active == 0, "the module should have deactivated the submission's own row"
    run = _query(
        config,
        "SELECT run_status, run_changes FROM actionrun WHERE project_id = '{}' "
        "AND form_id = '{}' AND main_rowuuid = '{}'".format(
            test_object.projectID, FORM, main2
        ),
    )
    assert run and run[0][0] == 0, run
    assert json.loads(run[0][1]) == [
        {
            "scope": "source",
            "table": "maintable",
            "rowuuid": main2,
            "column": "_active",
            "old": "1",
            "new": "0",
        }
    ]
    res = testapp.get(actions, status=200)
    test_object.root.assertIn(b"its rules act on its own rows", res.body)
