"""A case's properties in RSTools' Go exports (rstools.md 31 and 32).

The export plugins -- FastExports for CSV, JSON and XLSX, and the Stata one
-- run mysqltocsv2, mysqltojson2, mysqltoxlsx2 and mysqltostata2 with what
properties_argument gives them: ``-y`` and the form's properties.xml, so that
each table's properties go out as more columns beside its rows. The suite
loads neither plugin, so this runs the four exporters as they do:

- a form with no property adds nothing, and its export runs as before;
- with properties, properties.xml lists them, in create.xml's vocabulary;
- each format carries them after the table's own columns, typed, beside
  their rows, the properties' _lastupdate as _properties_lastupdate.

An exporter that predates -y is not given it (CI's image may lag), and the
step then stops after properties.xml.
"""

import json
import os
import shutil
import time
import uuid
import zipfile
from subprocess import PIPE, Popen

import openpyxl
from lxml import etree
from sqlalchemy import create_engine
from sqlalchemy.pool import NullPool

from formshare.processes.db.case_management import (
    export_properties_xml,
    exporter_takes_properties,
    properties_argument,
)
from .sql import get_form_details

FORM = "prop_export"
EXPORTERS = {
    "csv": ("MySQLToCSV2", "mysqltocsv2", "out.zip"),
    "json": ("MySQLToJSON2", "mysqltojson2", "out.zip"),
    "xlsx": ("MySQLToXLSX2", "mysqltoxlsx2", "out.xlsx"),
    "stata": ("MySQLToSTATA2", "mysqltostata2", "out.zip"),
}


def _query(config, sql):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        return engine.execute(sql).fetchall()
    finally:
        engine.dispose()


def _workbook(path):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "pid", "ID"])
    survey.append(["text", "farmer", "Farmer"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Properties in the exports", FORM, "1"])
    wb.save(path)


def _export(config, kind, schema, create_xml, work):
    """Runs one Go exporter as the plugins do, with properties_argument's
    arguments; returns the output's path."""
    folder, binary, name = EXPORTERS[kind]
    tool = os.path.join(config["odktools.path"], "utilities", folder, binary)
    out_dir = os.path.join(work, kind)
    os.makedirs(out_dir, exist_ok=True)
    output = os.path.join(out_dir, name)
    args = [
        tool,
        "-H=" + config["mysql.host"],
        "-P=" + str(config["mysql.port"]),
        "-u=" + config["mysql.user"],
        "-p=" + config["mysql.password"],
        "-s=" + schema,
        "-x=" + create_xml,
        "-o=" + output,
        "-e=8ee0f2c2cbba476db0a8a81fcde4cbee",
        "-w=2",
    ]
    extra, properties_file = properties_argument(
        config, schema, create_xml, tool, out_dir
    )
    assert extra == ["-y=" + properties_file], extra
    try:
        process = Popen(args + extra, stdout=PIPE, stderr=PIPE)
        stdout, stderr = process.communicate()
    finally:
        os.remove(properties_file)
    assert process.returncode == 0, (kind, stdout, stderr)
    return output


def t_e_s_t_properties_export(test_object):
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "properties_export")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)

    # --- A form, its repository and two farms -----------------------------------
    xlsx = os.path.join(work, FORM + ".xlsx")
    _workbook(xlsx)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "pid"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(login, project, FORM),
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )
    testapp.post(
        "/user/{}/project/{}/form/{}/repository/create".format(login, project, FORM),
        {"form_pkey": "pid", "start_stage1": ""},
        status=302,
    )
    schema = None
    for _ in range(100):
        schema = get_form_details(config, test_object.projectID, FORM)["form_schema"]
        if (
            schema
            and _query(
                config,
                "SELECT COUNT(*) FROM information_schema.TABLES WHERE "
                "TABLE_SCHEMA = '{}' AND TABLE_NAME = 'maintable'".format(schema),
            )[0][0]
        ):
            break
        time.sleep(3)
    assert schema, "the repository of {} did not build".format(FORM)
    for key, farmer in (("P1", "Wanjiru"), ("P2", "Otieno")):
        path = os.path.join(work, key + ".json")
        with open(path, "w") as a_file:
            json.dump(
                {
                    "_xform_id_string": FORM,
                    "pid": key,
                    "farmer": farmer,
                    "meta/instanceID": "uuid:{}".format(uuid.uuid4()),
                },
                a_file,
            )
        res = testapp.post(
            "/user/{}/project/{}/push_json".format(login, project),
            status=201,
            upload_files=[("filetoupload", path)],
            extra_environ=dict(
                FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin
            ),
        )
        assert "FS_error" not in res.headers
    create_xml = _query(
        config,
        "SELECT form_createxmlfile FROM odkform WHERE project_id = '{}' "
        "AND form_id = '{}'".format(test_object.projectID, FORM),
    )[0][0]

    # --- No property: nothing is added ------------------------------------------
    assert export_properties_xml(config, schema, create_xml) is None
    tool = os.path.join(
        config["odktools.path"], "utilities", "MySQLToCSV2", "mysqltocsv2"
    )
    assert properties_argument(config, schema, create_xml, tool, work) == ([], None)

    # --- Two properties of the farms --------------------------------------------
    # A property belongs on a table a published list draws from.
    res = testapp.post(
        "/user/{}/project/{}/caselists/add".format(login, project),
        {
            "list_id": "prop_farms",
            "list_filename": "prop_farms.csv",
            "source_form": FORM,
            "source_table": "maintable",
            "label_column": "farmer",
            "list_active": "1",
            "list_key_column": "",
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    properties = "/user/{}/project/{}/caseproperties?form={}&table=maintable".format(
        login, project, FORM
    )
    for name, ptype, default, desc in (
        ("status", "string", "open", "Where the farm is"),
        ("visits", "integer", "0", "Visits so far"),
    ):
        res = testapp.post(
            properties,
            {
                "add_property": "1",
                "property_name": name,
                "property_type": ptype,
                "property_default": default,
                "property_desc": desc,
            },
            status=302,
        )
        assert "FS_error" not in res.headers, name

    document = etree.fromstring(export_properties_xml(config, schema, create_xml))
    table = document.find("tables/table")
    assert table.get("name") == "maintable_properties"
    assert [f.get("name") for f in table.findall("field")] == [
        "rowuuid",
        "status",
        "visits",
        "_lastupdate",
    ]
    if not all(
        exporter_takes_properties(
            os.path.join(config["odktools.path"], "utilities", folder, binary)
        )
        for folder, binary, _ in EXPORTERS.values()
    ):
        return

    # --- Each format, the properties beside their rows ---------------------------
    expected_tail = ["status", "visits", "_properties_lastupdate"]
    output = _export(config, "csv", schema, create_xml, work)
    with zipfile.ZipFile(output) as a_zip:
        name = [n for n in a_zip.namelist() if n.endswith("maintable.csv")][0]
        lines = a_zip.read(name).decode("utf-8").splitlines()
    header = [h.strip('"') for h in lines[0].split(",")]
    assert header[-3:] == expected_tail, header
    assert len(lines) == 3
    for line in lines[1:]:
        values = [v.strip('"') for v in line.split(",")]
        assert values[-3:-1] == ["open", "0"], line

    output = _export(config, "json", schema, create_xml, work)
    with zipfile.ZipFile(output) as a_zip:
        rows = json.loads(a_zip.read("maintable.json"))
    assert len(rows) == 2
    for row in rows:
        assert list(row.keys())[-3:] == expected_tail, row
        assert (row["status"], row["visits"]) == ("open", 0), row

    output = _export(config, "xlsx", schema, create_xml, work)
    sheet = openpyxl.load_workbook(output).worksheets[0]
    header = [c.value for c in sheet[1]]
    assert header[-3:] == expected_tail, header
    for row in sheet.iter_rows(min_row=2, values_only=True):
        assert (row[-3], row[-2]) == ("open", 0), row

    output = _export(config, "stata", schema, create_xml, work)
    with zipfile.ZipFile(output) as a_zip:
        dta = a_zip.read("maintable.dta")
    for name in expected_tail:
        assert name.encode() in dta, name
