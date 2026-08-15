"""Offline entities: the parts that need no server.

`test_00_root` loads FormShare once and drives every route through it, which
takes well over an hour. Nothing here does that. These cover the pieces of the
offline entity work that are decidable without a database, a repository or an
ODK client, so the branch can be checked in seconds:

    export FORMSHARE_PYTEST_RUNNING=true
    pytest formshare/tests/test_01_entities.py -v

Naming an explicit file makes pytest ignore ``testpaths``, so the gate does not
run. They are cheap and load no environment, so they are also safe to leave in
a full run.

What is *not* here, because it needs the app: that a project actually opts in
through the form, that an upload actually reaches the injection, and that a
manifest served over HTTP carries the attribute. `steps/entities.py` holds that
flow, ready to be called from the gate.
"""

import os
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook
from pyxform import xls2xform

from formshare.processes.odk.api import (
    generate_manifest,
    set_rowuuid_from_entity,
)
from formshare.processes.odk.entities import (
    case_list_file_name,
    case_list_name,
    deployment_supports_entities,
    inject_entity_declaration,
    wrong_case_selector_file,
)
from formshare.views.odk import submission_response


class FakeRequest(object):
    """Enough of a request for the settings lookups these functions make."""

    def __init__(self, **settings):
        self.registry = mock.Mock()
        self.registry.settings = settings


def _creator_workbook(path, with_entities_sheet=False):
    """A case creator form as a user would upload one."""
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "hh_code", "Household code"])
    survey.append(["text", "hh_head", "Head of household"])
    survey.append(["date", "visit_date", "Visit date"])
    if with_entities_sheet:
        entities = workbook.create_sheet("entities")
        entities.append(["list_name", "label"])
        entities.append(["theirs", "${hh_head}"])
    settings = workbook.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Household registration", "hh_reg", "202608150001"])
    workbook.save(path)
    return path


class TestEntityCapability(unittest.TestCase):
    def test_off_unless_the_deployment_says_otherwise(self):
        """ODKTools cannot build an entity form, so silence means no."""
        self.assertFalse(deployment_supports_entities(FakeRequest()))
        self.assertFalse(
            deployment_supports_entities(
                FakeRequest(**{"odktools.supports_entities": "false"})
            )
        )

    def test_on_for_the_spellings_a_person_would_write(self):
        for value in ("true", "True", "TRUE", " yes ", "1"):
            self.assertTrue(
                deployment_supports_entities(
                    FakeRequest(**{"odktools.supports_entities": value})
                ),
                "{!r} should turn entities on".format(value),
            )


class TestCaseListName(unittest.TestCase):
    """The list name has to be settled before the creator form is uploaded, so
    it is derived from the project code rather than chosen."""

    def test_derives_from_the_project_code(self):
        self.assertEqual(case_list_name("demo"), "demo_cases")
        self.assertEqual(case_list_file_name("demo"), "demo_cases.csv")

    def test_is_a_usable_xml_identifier(self):
        # No dots, no spaces, no leading digit: the name is an XML identifier,
        # a file name and a secondary instance id all at once.
        for code in ("My Proj.01", "2024 test", "--", "ünïcode"):
            name = case_list_name(code)
            self.assertRegex(name, r"^[A-Za-z_][A-Za-z0-9_]*$", code)

    def test_is_stable(self):
        self.assertEqual(case_list_name("demo"), case_list_name("demo"))


class TestInjection(unittest.TestCase):
    """FormShare writes the entities sheet so the user does not have to."""

    def setUp(self):
        self.work = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), *["tmp", "entities"]
        )
        if not os.path.exists(self.work):
            os.makedirs(self.work)

    def test_pyxform_makes_an_offline_capable_form_from_it(self):
        """The whole point: 2024.1.0, which is what a client needs to create an
        entity while offline, and an id minted on the device."""
        source = _creator_workbook(os.path.join(self.work, "creator.xlsx"))
        injected, message = inject_entity_declaration(
            source, case_list_name("demo"), "hh_head"
        )
        self.assertTrue(injected, message)
        self.assertIn("entities", load_workbook(source).sheetnames)

        xml_file = os.path.join(self.work, "creator.xml")
        xls2xform.xls2xform_convert(source, xml_file)
        with open(xml_file, encoding="utf8") as form:
            xform = form.read()

        self.assertIn('entities:entities-version="2024.1.0"', xform)
        self.assertIn('dataset="demo_cases"', xform)
        self.assertIn('create="1"', xform)
        # The device mints the case identity, which is what makes an offline
        # registration referable before it has ever reached the server.
        self.assertIn("odk-instance-first-load", xform)
        self.assertIn("uuid()", xform)

    def test_leaves_a_workbook_that_already_declares_entities_alone(self):
        """A user who wrote their own sheet meant it, and a second declaration
        would collide. This is also what makes re-uploading a downloaded
        workbook safe, since the stored copy is the injected one."""
        source = _creator_workbook(
            os.path.join(self.work, "theirs.xlsx"), with_entities_sheet=True
        )
        injected, message = inject_entity_declaration(
            source, case_list_name("demo"), "hh_head"
        )
        self.assertTrue(injected, message)
        workbook = load_workbook(source)
        self.assertEqual(workbook.sheetnames.count("entities"), 1)
        self.assertEqual(workbook["entities"].cell(row=2, column=1).value, "theirs")

    def test_refuses_what_it_cannot_write(self):
        source = _creator_workbook(os.path.join(self.work, "legacy.xlsx"))
        legacy = os.path.join(self.work, "legacy.xls")
        os.replace(source, legacy)
        injected, message = inject_entity_declaration(legacy, "demo_cases", "hh_head")
        self.assertFalse(injected)
        self.assertIn("xlsx", message)

    def test_refuses_to_declare_a_case_it_cannot_label(self):
        source = _creator_workbook(os.path.join(self.work, "nolabel.xlsx"))
        injected, message = inject_entity_declaration(source, "demo_cases", "")
        self.assertFalse(injected)
        self.assertIn("label", message)


class TestSelectorFileCheck(unittest.TestCase):
    """A follow-up pointing at the wrong file resolves no list at all, and
    nothing reports it. This is the check that turns that into a sentence."""

    def setUp(self):
        self.request = FakeRequest(**{"odktools.supports_entities": "true"})
        self.patch = mock.patch(
            "formshare.processes.db.project.project_serves_entity_list",
            return_value=True,
        )
        self.serves = self.patch.start()
        self.addCleanup(self.patch.stop)

    def test_accepts_the_projects_own_case_list(self):
        self.assertIsNone(
            wrong_case_selector_file(self.request, "p1", "demo", "demo_cases.csv")
        )

    def test_names_the_file_the_form_should_have_used(self):
        self.assertEqual(
            wrong_case_selector_file(self.request, "p1", "demo", "households.csv"),
            "demo_cases.csv",
        )

    def test_leaves_a_barcode_alone(self):
        # A scan carries the case id rather than picking from a list. What the
        # code must contain is the rowuuid, which cannot be checked from here.
        self.assertIsNone(
            wrong_case_selector_file(self.request, "p1", "demo", "barcode")
        )

    def test_says_nothing_about_a_project_that_does_not_serve_entities(self):
        self.serves.return_value = False
        self.assertIsNone(
            wrong_case_selector_file(self.request, "p1", "demo", "households.csv")
        )

    def test_says_nothing_where_the_deployment_cannot_serve_entities(self):
        self.assertIsNone(
            wrong_case_selector_file(
                FakeRequest(), "p1", "demo", "households.csv"
            )
        )


class TestRowuuidFromEntity(unittest.TestCase):
    """The device mints the case identity; this is where it becomes the row's.

    The keys are the ones XMLtoJSON really produces -- an attribute is keyed by
    the path of its node, an @, and its name -- not the __id the entity model
    uses, which never appears in a submission.
    """

    UUID = "d4a4d3f6-1234-4abc-8def-0123456789ab"

    def test_a_created_entity_becomes_the_rows_identity(self):
        submission = {
            "hh_code": "HH-042",
            "meta/entity/@dataset": "demo_cases",
            "meta/entity/@create": "1",
            "meta/entity/@id": self.UUID,
        }
        set_rowuuid_from_entity(submission)
        self.assertEqual(submission["rowuuid"], self.UUID)

    def test_accepts_a_create_flag_written_as_a_word(self):
        submission = {"meta/entity/@id": self.UUID, "meta/entity/@create": "true"}
        set_rowuuid_from_entity(submission)
        self.assertEqual(submission["rowuuid"], self.UUID)

    def test_an_updated_entity_does_not(self):
        """On an update the attribute names the case being updated -- a row that
        already exists and whose rowuuid it already is. Taking it would make the
        follow-up claim the creator's identity and the insert would be refused."""
        submission = {"meta/entity/@id": self.UUID, "meta/entity/@update": "1"}
        set_rowuuid_from_entity(submission)
        self.assertNotIn("rowuuid", submission)

    def test_an_upsert_does_not_either(self):
        submission = {
            "meta/entity/@id": self.UUID,
            "meta/entity/@create": "1",
            "meta/entity/@update": "1",
        }
        set_rowuuid_from_entity(submission)
        self.assertNotIn("rowuuid", submission)

    def test_an_ordinary_submission_is_untouched(self):
        submission = {"hh_code": "HH-042", "meta/instanceID": "uuid:1111"}
        set_rowuuid_from_entity(submission)
        self.assertNotIn("rowuuid", submission)
        self.assertEqual(len(submission), 2)

    def test_nothing_is_claimed_without_an_id(self):
        for submission in (
            {"meta/entity/@id": "", "meta/entity/@create": "1"},
            {"meta/entity/@id": self.UUID, "meta/entity/@create": "0"},
            {"meta/entity/@id": self.UUID},
        ):
            set_rowuuid_from_entity(submission)
            self.assertNotIn("rowuuid", submission, submission)


class TestManifest(unittest.TestCase):
    """A client tells an entity list from an attachment by an attribute, and the
    manifest builder could only write child elements."""

    def _media(self, **extra):
        media = {
            "filename": "demo_cases.csv",
            "hash": "md5:9a0f",
            "downloadUrl": "https://example.org/demo_cases.csv",
        }
        media.update(extra)
        return media

    def test_type_is_an_attribute_not_an_element(self):
        xml = generate_manifest([self._media(type="entityList")]).decode()
        self.assertIn('<mediaFile type="entityList">', xml)
        self.assertNotIn("<type>", xml)

    def test_everything_else_is_still_a_child(self):
        xml = generate_manifest(
            [self._media(type="entityList", integrityUrl="https://example.org/i")]
        ).decode()
        self.assertIn("<filename>demo_cases.csv</filename>", xml)
        self.assertIn("<integrityUrl>https://example.org/i</integrityUrl>", xml)

    def test_an_ordinary_media_file_gains_nothing(self):
        """A project not serving entities must produce the manifest it always
        did, or every existing form's attachments change shape."""
        for media in (self._media(), self._media(type=None)):
            xml = generate_manifest([media]).decode()
            self.assertIn("<mediaFile>", xml)
            self.assertNotIn("type=", xml)


class TestSubmissionResponse(unittest.TestCase):
    """A submission that did not reach the repository is still accepted, and is
    waiting in the logs. Saying so is the difference between an enumerator
    knowing and an enumerator walking away."""

    def test_says_nothing_when_there_is_nothing_to_say(self):
        response = submission_response(201)
        self.assertEqual(response.status, 201)
        self.assertEqual(response.body, b"")

    def test_carries_the_message_without_changing_the_status(self):
        response = submission_response(201, "Sent to the logs.")
        self.assertEqual(response.status, 201)
        body = response.body.decode()
        self.assertIn("<OpenRosaResponse", body)
        self.assertIn("http://openrosa.org/http/response", body)
        self.assertIn("<message>Sent to the logs.</message>", body)
        self.assertEqual(response.headers.get("X-OpenRosa-Version"), "1.0")

    def test_a_message_cannot_break_the_xml(self):
        response = submission_response(201, 'Case "A&B" <x> failed')
        body = response.body.decode()
        self.assertIn("&amp;", body)
        self.assertIn("&lt;x&gt;", body)
        from xml.dom.minidom import parseString

        parseString(body)


if __name__ == "__main__":
    unittest.main()
