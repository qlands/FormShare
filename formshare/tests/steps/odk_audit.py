"""A submission's ODK audit, from the device to its table
(processes/odk_audit.py, docs/odk-audit.md).

A form with an audit row and track-changes, in a project with a key, that
encrypts farmer_name and the repeat's mem_name: with track-changes, Collect
writes those answers into audit.csv in clear. The audits are written here in
Collect's own format (AuditEventSaveTask), since no device sent one to a
dev server, and every answer carries this run's marker, so that looking for
it in the files cannot be fooled by an earlier run.

- the build makes odk_audit, and the unique index on maintable.surveyid its
  key needs;
- a submission's audit gets its rows: the marked values sealed, in the rows
  and in the file kept, the others as written;
- a resend changes nothing; an audit in the last part of a submission sent
  in parts goes to the first part's submission;
- a submission that goes to the logs gets no rows; a deleted one loses them;
  the one pushed back from the logs gets them;
- an import loads the audit of what it imports, sealing a value in clear;
- a form in testing keeps its audit sealed; an audit of such a form that
  cannot be read is not kept;
- deleting all submissions empties the table;
- a repository built before the table gets it at its next merge;
- a repeat called odk_audit is refused at upload.

The repository is built and the version merged in this process, as the
worker does (in_process), so no Celery worker's code matters here.
"""

import glob
import json
import os
import shutil
import time
import uuid

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import NullPool

from formshare.processes.db.encryption import (
    enroll_member,
    get_member,
    is_envelope,
    member_name,
    set_up_project_encryption,
)
from formshare.processes.odk_audit import (
    AUDIT_FILE,
    AUDIT_TABLE,
    load_audit,
    make_audit_table,
    make_audit_table_for,
)
from formshare.products.fs1import.celery_task import internal_import_json_files
from .encryption_forms import (
    _assistant,
    _give,
    _new_project,
    _new_version,
    _project_id,
    _query,
    _upload,
)
from .encryption_keys import _Person, _Request, _grant, _project_key, _signed
from .encryption_submissions import _nowhere, _open, _since
from .in_process import add_task, build_in_process, form_row
from .merge_manifest import SURVEY_COLUMNS
from .sql import get_form_details

CODE = "odk_audit"
FORM = "aud_form"
HEADER = ["type", "name", "label", "fs_encrypted", "parameters"]
AUDITED_FORM = [
    ["text", "pid", "ID", "", ""],
    ["text", "farmer_name", "Farmer", "yes", ""],
    ["integer", "age", "Age", "", ""],
    ["text", "notes", "Notes", "", ""],
    ["begin repeat", "member", "Member", "", ""],
    ["text", "mem_name", "Name", "yes", ""],
    ["end repeat", "", "", "", ""],
    ["audit", "audit", "", "", "track-changes=true"],
]
# What Collect writes with track-changes and location (AuditEventSaveTask)
AUDIT_HEADER = "event,node,start,end,latitude,longitude,accuracy,old-value,new-value"


def _workbook(path, form_id, rows, version="1"):
    book = openpyxl.Workbook()
    survey = book.active
    survey.title = "survey"
    survey.append(HEADER)
    for a_row in rows:
        survey.append(a_row)
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append([form_id, form_id, version])
    book.save(path)
    return path


def _xml(path, form_id, pid, farmer, member, version="1"):
    with open(path, "w") as a_file:
        a_file.write(
            '<?xml version="1.0"?>\n<data id="{}" version="{}"><pid>{}</pid>'
            "<farmer_name>{}</farmer_name><age>34</age><notes>kept in clear</notes>"
            "<member><mem_name>{}</mem_name></member><meta><audit>audit.csv</audit>"
            "<instanceID>uuid:{}</instanceID></meta></data>".format(
                form_id, version, pid, farmer, member, uuid.uuid4()
            )
        )
    return path


def _audit(work, name, pid, farmer, member, changed_to):
    """An audit.csv as Collect writes it for one interview: the farmer's name
    typed, then changed after the repeat, so it has an old and a new value."""
    directory = os.path.join(work, name)
    os.makedirs(directory)
    path = os.path.join(directory, "audit.csv")
    start = 1761800000000
    where = "-1.2921,36.8219,5.0"
    lines = [
        AUDIT_HEADER,
        "form start,,{},,,,,,".format(start),
        "question,/data/pid,{},{},{},,{}".format(start + 100, start + 2100, where, pid),
        'question,/data/farmer_name,{},{},{},,"{}"'.format(
            start + 2100, start + 9000, where, farmer
        ),
        'question,/data/notes,{},{},{},,"kept, in clear"'.format(
            start + 9000, start + 12000, where
        ),
        "add repeat,/data/member[1],{},{},{},,".format(
            start + 12000, start + 13000, where
        ),
        "question,/data/member[1]/mem_name,{},{},{},,{}".format(
            start + 13000, start + 15000, where, member
        ),
        'question,/data/farmer_name,{},{},{},"{}",{}'.format(
            start + 15000, start + 17000, where, farmer, changed_to
        ),
        "end screen,,{},{},{},,".format(start + 17000, start + 18000, where),
        "form save,,{},,{},,".format(start + 18000, where),
        "form exit,,{},,{},,".format(start + 18000, where),
        "form finalize,,{},,{},,".format(start + 18000, where),
    ]
    with open(path, "w", newline="") as a_file:
        a_file.write("\n".join(lines) + "\n")
    return path


def _push(testapp, login, path, assistant, media=(), status=201, incomplete=False):
    return testapp.post(
        "/user/{}/project/{}/push".format(login, CODE),
        {"*isIncomplete*": "yes"} if incomplete else {},
        status=status,
        upload_files=[("xml_submission_file", path)]
        + [(os.path.basename(a_file), a_file) for a_file in media],
        extra_environ=dict(FS_for_testing="true", FS_user_for_testing=assistant[0]),
    )


def _submission_of(engine, schema, pid):
    rows = _query(
        engine,
        "SELECT surveyid, rowuuid FROM `{}`.maintable WHERE pid = %(pid)s".format(
            schema
        ),
        pid=pid,
    )
    return rows[0] if rows else (None, None)


def _rows(engine, schema, submission_id):
    return _query(
        engine,
        "SELECT event_order, event, node, start, `end`, latitude, longitude, "
        "accuracy, old_value, new_value FROM `{}`.`odk_audit` "
        "WHERE submission_id = %(s)s ORDER BY event_order".format(schema),
        s=submission_id,
    )


def _count(engine, schema):
    return _query(engine, "SELECT COUNT(*) FROM `{}`.`odk_audit`".format(schema))[0][0]


def _has_table(engine, schema):
    return bool(
        _query(
            engine,
            "SELECT COUNT(*) FROM information_schema.tables WHERE "
            "table_schema = %(s)s AND table_name = 'odk_audit'",
            s=schema,
        )[0][0]
    )


def _keyed(engine, schema):
    """The table's key to maintable, cascading, and maintable's unique index
    on surveyid that the key needs."""
    cascade = _query(
        engine,
        "SELECT delete_rule FROM information_schema.referential_constraints "
        "WHERE constraint_schema = %(s)s AND table_name = 'odk_audit' "
        "AND referenced_table_name = 'maintable'",
        s=schema,
    )
    unique = _query(
        engine,
        "SELECT non_unique FROM information_schema.statistics WHERE "
        "table_schema = %(s)s AND table_name = 'maintable' AND "
        "index_name = 'fs_surveyid'",
        s=schema,
    )
    return [row[0] for row in cascade] == ["CASCADE"] and [
        row[0] for row in unique
    ] == [0]


def _set_up_key(engine, session, login, project_id):
    """The project's key, set up as the encryption steps set it up. Returns
    its private half, to open what was sealed."""
    request = _Request(session)
    existing = get_member(request, "user", login)
    if existing is None:
        owner = _Person()
        ok, owner_id = enroll_member(
            request, "user", login, owner.public_key, owner.unlock()
        )
        assert ok, owner_id
        owner_public = owner.public_key
    else:
        owner_id, owner_public = existing["member_id"], existing["public_key"]
    project_private, public_key = _project_key()
    recovery = _Person()
    ok, key_id = set_up_project_encryption(
        request,
        project_id,
        public_key,
        owner_id,
        _grant(project_private, member_name("user", login), owner_public),
        recovery.public_key,
        _grant(
            project_private,
            member_name("recovery", project_id),
            recovery.public_key,
        ),
        login,
        **_signed(project_private),
    )
    assert ok, key_id
    return project_private


def _merge(test_object, project_id, version, parent):
    """The merge check as the version's page opens it, then the merge in this
    process, as the worker runs it (merge_manifest)."""
    from formshare.products.merge.celery_task import internal_merge_into_repository

    login = test_object.randonLogin
    config = test_object.server_config
    testapp = test_object.testapp
    res = testapp.get(
        "/user/{}/project/{}/form/{}".format(login, CODE, version), status=200
    )
    assert b"Merge check pending" not in res.body
    testapp.get(
        "/user/{}/project/{}/form/{}/merge/into/{}".format(
            login, CODE, version, parent
        ),
        status=200,
    )
    a_form = form_row(config, project_id, version)
    b_form = form_row(config, project_id, parent)
    internal_merge_into_repository(
        config,
        login,
        project_id,
        CODE,
        version,
        a_form["form_directory"],
        b_form["form_directory"],
        b_form["form_createxmlfile"],
        b_form["form_schema"],
        "",
        b_form["form_hexcolor"],
        "en",
        True,
        SURVEY_COLUMNS,
        add_task(config, project_id, version, "merge_form", login),
    )
    a_form = form_row(config, project_id, version)
    assert a_form["form_schema"] == b_form["form_schema"]
    assert a_form["form_blocked"] == 0


def t_e_s_t_odk_audit(test_object):
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "odk_audit")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    session = sessionmaker(bind=engine)()
    try:
        _run(test_object, work, engine, session)
    finally:
        session.close()
        engine.dispose()


def _run(test_object, work, engine, session):
    login = test_object.randonLogin
    testapp = test_object.testapp
    config = test_object.server_config
    marker = uuid.uuid4().hex[:8]
    started = time.time()
    odk = os.path.join(config["repository.path"], "odk")
    staging = os.path.join(odk, "submissions")

    _new_project(testapp, login, CODE)
    project_id = _project_id(engine, login, CODE)
    assistant = _assistant(testapp, login, CODE)
    project_private = _set_up_key(engine, session, login, project_id)

    def opened(value):
        assert is_envelope(value), value
        return _open(value, project_private)

    # --- A repeat called odk_audit is refused at upload ---------------------------
    taken = _workbook(
        os.path.join(work, "aud_taken.xlsx"),
        "aud_taken",
        [
            ["text", "pid", "ID", "", ""],
            ["begin repeat", "odk_audit", "Audit", "", ""],
            ["text", "auditor", "Auditor", "", ""],
            ["end repeat", "", "", "", ""],
        ],
    )
    res = _upload(testapp, login, CODE, taken)
    assert "FS_error" in res.headers
    res = res.follow(status=200)
    assert b"The form has a repeat called" in res.body
    assert (
        _query(
            engine,
            "SELECT COUNT(*) FROM odkform WHERE project_id = %(p)s AND "
            "form_id = 'aud_taken'",
            p=project_id,
        )[0][0]
        == 0
    )

    # --- The build makes the table, keyed to maintable ---------------------------
    for form_id in (FORM, "aud_testing"):
        res = _upload(
            testapp,
            login,
            CODE,
            _workbook(os.path.join(work, form_id + ".xlsx"), form_id, AUDITED_FORM),
        )
        assert "FS_error" not in res.headers
        _give(testapp, login, CODE, form_id, project_id, assistant)
    build_in_process(test_object, project_id, CODE, FORM, "pid", [])
    details = get_form_details(config, project_id, FORM)
    schema = details["form_schema"]
    assert schema
    assert _has_table(engine, schema)
    assert _keyed(engine, schema)
    form_directory = os.path.join(odk, "forms", details["form_directory"])

    # --- A submission's audit: its rows, the marked values sealed -------------------
    kamau = "Kamau, " + marker
    changed = "Kamau Mwangi " + marker
    njeri = "Njeri " + marker
    s1_audit = _audit(work, "s1", "S1", kamau, njeri, changed)
    s1_xml = _xml(os.path.join(work, "s1.xml"), FORM, "S1", changed, njeri)
    _push(testapp, login, s1_xml, assistant, [s1_audit])
    s1, s1_rowuuid = _submission_of(engine, schema, "S1")
    assert s1 is not None
    rows = _rows(engine, schema, s1)
    assert [row[0] for row in rows] == list(range(1, 12))
    assert rows[0][1:4] == ("form start", None, 1761800000000)
    assert rows[0][4] is None
    pid_row = rows[1]
    assert pid_row[1:3] == ("question", "/data/pid")
    assert (pid_row[3], pid_row[4]) == (1761800000100, 1761800002100)
    assert (pid_row[5], pid_row[6], pid_row[7]) == (-1.2921, 36.8219, 5.0)
    assert pid_row[8] is None and pid_row[9] == "S1"
    assert opened(rows[2][9]) == kamau
    assert rows[3][8:10] == (None, "kept, in clear")
    assert rows[4][1:3] == ("add repeat", "/data/member[1]")
    assert opened(rows[5][9]) == njeri
    assert opened(rows[6][8]) == kamau and opened(rows[6][9]) == changed
    assert [row[1] for row in rows[7:]] == [
        "end screen",
        "form save",
        "form exit",
        "form finalize",
    ]
    kept = os.path.join(form_directory, "submissions", s1, "audit.csv")
    with open(kept) as a_file:
        content = a_file.read()
    assert "kept, in clear" in content and content.startswith(AUDIT_HEADER)

    # --- A resend changes nothing --------------------------------------------------
    _push(testapp, login, s1_xml, assistant, [s1_audit])
    assert _rows(engine, schema, s1) == rows
    assert _count(engine, schema) == 11

    # --- An audit in the last part goes to the first part's submission ----------------
    wanjiru = "Wanjiru " + marker
    otieno = "Otieno " + marker
    s2_xml = _xml(os.path.join(work, "s2.xml"), FORM, "S2", wanjiru, otieno)
    photo = os.path.join(work, "part1.jpg")
    with open(photo, "wb") as a_file:
        a_file.write(os.urandom(2048))
    _push(testapp, login, s2_xml, assistant, [photo], incomplete=True)
    s2, _rowuuid = _submission_of(engine, schema, "S2")
    assert _rows(engine, schema, s2) == []
    _push(
        testapp,
        login,
        s2_xml,
        assistant,
        [_audit(work, "s2", "S2", wanjiru, otieno, wanjiru + " K")],
    )
    rows = _rows(engine, schema, s2)
    assert len(rows) == 11 and opened(rows[2][9]) == wanjiru
    s2_media = os.path.join(form_directory, "submissions", s2)
    assert sorted(os.listdir(s2_media)) == ["audit.csv", "diffs", "part1.jpg"]

    # --- The logs: no rows there, none for a deleted one, rows when it comes back -----
    achieng = "Achieng " + marker
    s3_xml = _xml(os.path.join(work, "s3.xml"), FORM, "S1", achieng, "Odhiambo")
    _push(
        testapp,
        login,
        s3_xml,
        assistant,
        [_audit(work, "s3", "S1", achieng, "Odhiambo", achieng + " A")],
    )
    in_logs = _query(
        engine,
        "SELECT submission_id FROM submission WHERE project_id = %(p)s AND "
        "form_id = %(f)s AND submission_status = 2 AND sameas IS NULL",
        p=project_id,
        f=FORM,
    )
    assert len(in_logs) == 1, in_logs
    s3 = in_logs[0][0]
    assert _rows(engine, schema, s3) == []
    res = testapp.post(
        "/user/{}/project/{}/form/{}/submissions/delete".format(login, CODE, FORM),
        {"oper": "del", "id": s1_rowuuid},
        status=200,
    )
    assert _rows(engine, schema, s1) == []
    res = testapp.post(
        "/user/{}/project/{}/assistantaccess/login".format(login, CODE),
        {"login": assistant[0], "passwd": "123"},
        status=302,
    )
    assert "FS_error" not in res.headers
    res = testapp.post(
        "/user/{}/project/{}/assistantaccess/form/{}/{}/push".format(
            login, CODE, FORM, s3
        ),
        status=302,
    )
    assert "FS_error" not in res.headers
    testapp.post(
        "/user/{}/project/{}/assistantaccess/logout".format(login, CODE), status=302
    )
    rows = _rows(engine, schema, s3)
    assert len(rows) == 11 and opened(rows[2][9]) == achieng

    # --- Nothing answered in clear is on the disk ------------------------------------
    _nowhere(
        [form_directory] + _since(staging, started),
        kamau,
        changed,
        njeri,
        wanjiru,
        otieno,
        achieng,
    )

    # --- An import loads the audit of what it imports, sealing what is clear --------
    to_import = os.path.join(work, "to_import")
    os.makedirs(to_import)
    s4 = str(uuid.uuid4())
    with open(os.path.join(to_import, s4 + ".json"), "w") as a_file:
        json.dump(
            {
                "_xform_id_string": FORM,
                "pid": "S4",
                "farmer_name": "Imported " + marker,
                "notes": "kept in clear",
            },
            a_file,
        )
    # As a file kept before FormShare sealed audits would be
    planted = "Planted " + marker
    os.makedirs(os.path.join(form_directory, "submissions", s4))
    shutil.copy(
        _audit(work, "s4", "S4", planted, "Nobody", planted + " B"),
        os.path.join(form_directory, "submissions", s4, "audit.csv"),
    )
    internal_import_json_files(
        login,
        project_id,
        FORM,
        odk,
        details["form_directory"],
        schema,
        assistant[1],
        to_import,
        CODE,
        [],
        config,
        "en",
        report_task=False,
    )
    rows = _rows(engine, schema, s4)
    assert len(rows) == 11
    assert opened(rows[2][9]) == planted and opened(rows[6][9]) == planted + " B"
    assert rows[3][9] == "kept, in clear"

    # --- A form in testing keeps its audit sealed ----------------------------------------
    testing_directory = os.path.join(
        odk,
        "forms",
        get_form_details(config, project_id, "aud_testing")["form_directory"],
    )
    mwangi = "Mwangi " + marker
    _push(
        testapp,
        login,
        _xml(os.path.join(work, "t1.xml"), "aud_testing", "T1", mwangi, "Akinyi"),
        assistant,
        [_audit(work, "t1", "T1", mwangi, "Akinyi", mwangi + " M")],
    )
    kept = glob.glob(os.path.join(testing_directory, "submissions", "*", "audit.csv"))
    assert len(kept) == 1, kept
    with open(kept[0]) as a_file:
        content = a_file.read()
    assert "kept, in clear" in content
    _nowhere([testing_directory] + _since(staging, started), mwangi)

    # --- An audit that cannot be read is not kept for a form that encrypts -----------------
    unreadable = os.path.join(work, "s5")
    os.makedirs(unreadable)
    with open(os.path.join(unreadable, "audit.csv"), "w") as a_file:
        a_file.write("name,answer\nfarmer_name,Hidden {}\n".format(marker))
    _push(
        testapp,
        login,
        _xml(os.path.join(work, "s5.xml"), FORM, "S5", "Hidden " + marker, "x"),
        assistant,
        [os.path.join(unreadable, "audit.csv")],
    )
    s5, _rowuuid = _submission_of(engine, schema, "S5")
    assert s5 is not None
    assert not os.path.exists(
        os.path.join(form_directory, "submissions", s5, AUDIT_FILE)
    )
    assert _rows(engine, schema, s5) == []
    _nowhere([form_directory] + _since(staging, started), "Hidden " + marker)

    # --- Deleting all submissions empties the table -----------------------------------------
    res = testapp.post(
        "/user/{}/project/{}/form/{}/submissions/deleteall".format(login, CODE, FORM),
        {"owner_email": login + "@qlands.com"},
        status=302,
    )
    assert "FS_error" not in res.headers
    assert _has_table(engine, schema) and _count(engine, schema) == 0

    # --- A repository built before the table gets it at its next merge --------------------
    with engine.connect() as connection:
        connection.execute("DROP TABLE `{}`.`{}`".format(schema, AUDIT_TABLE))
        connection.execute(
            "ALTER TABLE `{}`.maintable DROP INDEX fs_surveyid".format(schema)
        )
    assert not _has_table(engine, schema)
    _new_version(
        testapp,
        login,
        CODE,
        FORM,
        project_id,
        _workbook(
            os.path.join(work, "aud_form_b.xlsx"),
            "aud_form_b",
            AUDITED_FORM[:4]
            + [["text", "village", "Village", "", ""]]
            + AUDITED_FORM[4:],
            version="2",
        ),
    )
    _merge(test_object, project_id, "aud_form_b", FORM)
    assert _has_table(engine, schema)
    assert _keyed(engine, schema)
    # The next merge finds it there and leaves it as it is
    assert make_audit_table(engine, schema) == (True, "")

    # --- Where the table cannot be made, and what is then not loaded --------------------
    # A loaded submission's row exists; one that is not in the repository
    # gets nothing
    assert load_audit(engine, schema, "not-a-submission", s1_audit) is None
    # Nor does a form that encrypts, without the project's key
    assert load_audit(engine, schema, s2, s1_audit, {"farmer_name"}, None) is None
    scratch = "FS_odk_audit_" + marker
    with engine.connect() as connection:
        connection.execute("CREATE SCHEMA `{}`".format(scratch))
    try:
        assert make_audit_table(engine, scratch) == (
            False,
            "the repository has no maintable",
        )
        with engine.connect() as connection:
            connection.execute(
                "CREATE TABLE `{}`.maintable (pid VARCHAR(10) PRIMARY KEY) "
                "ENGINE=InnoDB".format(scratch)
            )
        assert make_audit_table(engine, scratch) == (
            False,
            "maintable has no surveyid",
        )
        with engine.connect() as connection:
            connection.execute(
                "ALTER TABLE `{}`.maintable ADD COLUMN surveyid VARCHAR(80)".format(
                    scratch
                )
            )
            connection.execute(
                "INSERT INTO `{}`.maintable VALUES ('A', 'twice'), ('B', 'twice')".format(
                    scratch
                )
            )
        assert make_audit_table(engine, scratch) == (
            False,
            "two submissions have the ID twice",
        )
        # A repository without the table loads nothing
        assert load_audit(engine, scratch, "twice", s1_audit) is None
        with engine.connect() as connection:
            connection.execute(
                "CREATE TABLE `{}`.`{}` (auditor VARCHAR(10))".format(
                    scratch, AUDIT_TABLE
                )
            )
        assert make_audit_table(engine, scratch) == (
            False,
            "a table of the form is called odk_audit",
        )
        # A build or a merge goes on without it, and says why
        assert make_audit_table_for(config, scratch) is False
        assert load_audit(engine, scratch, "twice", s1_audit) is None
    finally:
        with engine.connect() as connection:
            connection.execute("DROP SCHEMA IF EXISTS `{}`".format(scratch))
