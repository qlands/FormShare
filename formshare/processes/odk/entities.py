"""
formshare.processes.odk.entities
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~

Serving a project's case list to ODK clients as an **entity list**, so a case
registered offline can be followed up before the device has synced.

FormShare has held cases relationally since 2021: a case is a row in the case
creator form's repository and a trigger refuses a follow-up whose creator row is
not active. None of that changes here. What changes is that the case list also
travels as an ODK entity list, which the client keeps in its own store and can
add to while offline.

The user is not asked to write an ``entities`` sheet. FormShare already knows
which variable labels a case, so it writes the sheet itself into the uploaded
workbook before the workbook reaches pyxform. pyxform does the rest: it emits
``entities-version="2024.1.0"``, mints the entity id on the device, and builds
the ``<entity>`` block.

Two things are deliberately narrow:

* **Only the case creator form is injected.** Follow-up forms consume the list
  through the ``select_one_from_file`` they already declare, and do not update
  entities. Whether a follow-up may modify the case is not settled.
* **No ``save_to`` properties are written.** The list of properties comes from
  ``CaseLookUp``, which the user fills in *after* the creator form is uploaded
  and its repository built, so there is nothing to read at injection time. An
  offline created case therefore carries its label and nothing else until it
  syncs. A follow-up form that filters the list on a property will not match a
  case created offline until then.
"""

import logging
import os
import re

log = logging.getLogger("formshare")

__all__ = [
    "deployment_supports_entities",
    "project_uses_entities",
    "case_list_name",
    "case_list_file_name",
    "wrong_case_selector_file",
    "inject_entity_declaration",
]

# The dataset name is an XML identifier: no dots, no leading __, and it has to
# survive being used as a file name and as a secondary instance id.
_INVALID_IN_NAME = re.compile(r"[^A-Za-z0-9_]")


def deployment_supports_entities(request):
    """Whether the tool chain installed here can build an entity carrying form.

    RSTools skips the entity declaration when generating the schema and mints v4
    rowuuids. ODKTools, the lighter build, does neither: it would turn the
    declaration into a stray ``entity`` column and mint v1 ids that the client
    silently refuses. The same FormShare code runs against both, so the
    deployment has to say which one it has.
    """
    setting = request.registry.settings.get("odktools.supports_entities", "false")
    return str(setting).strip().lower() in ("true", "1", "yes")


def project_uses_entities(request, project_details):
    """Whether this project's case list should be served as an entity list.

    Both gates have to be open: the deployment can do it, and the project asked
    for it. Projects that predate the feature have ``project_entities`` at 0 and
    keep the relational-only behaviour they were built with -- their
    repositories still carry the old v1 rowuuid trigger.

    :param project_details: a project dict, as returned by get_project_details
    """
    if not deployment_supports_entities(request):
        return False
    if not project_details:
        return False
    if int(project_details.get("project_case", 0) or 0) != 1:
        return False
    return int(project_details.get("project_entities", 0) or 0) == 1


def case_list_name(project_code):
    """The name of a project's entity list, fixed for the life of the project.

    It has to be known when the creator form is uploaded, because that is when
    the ``<entity dataset>`` is written -- but the file name the client asks for
    comes from a *follow-up* form, which does not exist yet. Deriving it from the
    project code is what makes both sides agree without anyone choosing twice.

    The client names the list after the media file minus its extension, so this
    same name is what a follow-up form must reference in
    ``select_one_from_file <name>.csv``.
    """
    name = _INVALID_IN_NAME.sub("_", str(project_code or "").strip())
    name = name.lstrip("_") or "cases"
    if name[0].isdigit():
        name = "c" + name
    return "{}_cases".format(name.lower())


def case_list_file_name(project_code):
    """The media file name a follow-up form has to reference."""
    return case_list_name(project_code) + ".csv"


def wrong_case_selector_file(request, project_id, project_code, selector_file):
    """Checks the file a follow-up form selects cases from.

    A client names an entity list after the media file it arrived in, minus the
    extension, and that name has to be the one the creator form declared as its
    ``<entity dataset>``. FormShare fixes that name per project so both sides can
    know it without anyone choosing twice -- but the follow-up form is written by
    hand, so it can point somewhere else, and if it does the form simply never
    resolves the list. Nothing errors; the select is just empty.

    Catching it at upload turns a silent nothing into a sentence.

    :return: the file name that was expected, or None when there is nothing
        wrong -- the project does not serve entities, or the form already agrees
    """
    from formshare.processes.db.project import project_serves_entity_list

    if not deployment_supports_entities(request):
        return None
    if not project_serves_entity_list(request, project_id):
        return None
    if selector_file == "barcode":
        # A scan carries the case id itself rather than picking from the list,
        # so there is no file to agree about. What the code has to contain is
        # the rowuuid, which cannot be checked from here.
        return None
    expected = case_list_file_name(project_code)
    if selector_file == expected:
        return None
    return expected


def inject_entity_declaration(xlsx_file, dataset, label_variable):
    """Writes an ``entities`` sheet into an uploaded workbook, in place.

    Called between the upload landing on disk and pyxform reading it, so that
    one edit feeds both the XForm the device downloads and the JSON the schema
    is built from -- they cannot drift apart.

    :param xlsx_file: the uploaded workbook, modified in place
    :param dataset: the entity list name, from case_list_name()
    :param label_variable: the survey variable that labels a case, which the
        project owner named at upload as form_caselabel
    :return: (True, "") or (False, reason)
    """
    from openpyxl import load_workbook

    if os.path.splitext(xlsx_file)[1].lower() not in (".xlsx", ".xlsm"):
        # .xls cannot be written back out; it is legacy and unsupported here.
        return False, "Entity support needs an .xlsx or .xlsm file"

    if not label_variable:
        return False, "Cannot declare an entity without the case label variable"

    try:
        workbook = load_workbook(xlsx_file)
    except Exception as e:
        log.error("Cannot open {} to declare an entity. Error: {}".format(xlsx_file, e))
        return False, str(e)

    # A workbook that already declares entities is left alone: the user meant it,
    # and a second declaration would collide with theirs.
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() == "entities":
            log.info(
                "{} already has an entities sheet. Leaving it alone.".format(xlsx_file)
            )
            return True, ""

    try:
        sheet = workbook.create_sheet("entities")
        sheet.append(["list_name", "label"])
        sheet.append([dataset, "${" + label_variable + "}"])
        workbook.save(xlsx_file)
    except Exception as e:
        log.error(
            "Cannot write the entities sheet into {}. Error: {}".format(xlsx_file, e)
        )
        return False, str(e)

    log.info(
        "Declared entity list {} in {} labelled by {}".format(
            dataset, xlsx_file, label_variable
        )
    )
    return True, ""
