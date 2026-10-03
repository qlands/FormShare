"""Exporting and importing published lists, from one project into another.

A training manual's example is built in one project and loaded into another:
the same form, uploaded there, has a schema of its own, and the imported
lists must follow it. Three lists are built through the wizard in a project
of their own -- a list of rows with a served column, a property and a
filter, a value list, a list from a repeat -- exported one at a time and all
together, as JSON and as YAML, and imported into a second project:

- refused while that project lacks the form, then while the form lacks its
  repository, then when a list names a column the table does not have, with
  nothing written each time;
- accepted once everything is in place: the lists point at the second
  project's form, the property is created in its repository, the filter is
  compiled again for it, and its export is the first project's export;
- refused the second time, every code being in use.
"""

import copy
import json
import os
import shutil
import time
import uuid

import openpyxl
import yaml
from sqlalchemy import create_engine
from sqlalchemy.pool import NullPool

from .sql import get_form_details

FORM = "transfer_reg"
SOURCE = "listsource"
TARGET = "listimport"


def _workbook(path):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "farmer_id", "Farmer ID"])
    survey.append(["text", "farmer_name", "Farmer name"])
    survey.append(["select_one county", "county", "County"])
    survey.append(["begin_repeat", "members", "Members"])
    survey.append(["text", "member_name", "Member name"])
    survey.append(["end_repeat", "", ""])
    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    choices.append(["county", "kitui", "Kitui"])
    choices.append(["county", "laikipia", "Laikipia"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Farmers for the list transfer", FORM, "1"])
    wb.save(path)


def _query(config, sql):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        return engine.execute(sql).fetchall()
    finally:
        engine.dispose()


def _new_project(testapp, config, login, code):
    res = testapp.post(
        "/user/{}/projects/add".format(login),
        {
            "project_id": str(uuid.uuid4()),
            "project_code": code,
            "project_name": "List transfer " + code,
            "project_abstract": "",
            "project_icon": "",
            "project_hexcolor": "",
            "project_formlist_auth": 1,
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    return _query(
        config,
        "SELECT p.project_id FROM project p JOIN userproject u "
        "ON u.project_id = p.project_id WHERE u.user_id = '{}' "
        "AND p.project_code = '{}'".format(login, code),
    )[0][0]


def _upload(testapp, login, code, xlsx):
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, code),
        {"form_pkey": "farmer_id"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers


def _build(testapp, config, login, code, project_id):
    testapp.post(
        "/user/{}/project/{}/form/{}/repository/create".format(login, code, FORM),
        {"form_pkey": "farmer_id", "start_stage1": ""},
        status=302,
    )
    for _ in range(100):
        schema = get_form_details(config, project_id, FORM)["form_schema"]
        if (
            schema
            and _query(
                config,
                "SELECT COUNT(*) FROM information_schema.TABLES WHERE TABLE_SCHEMA = '{}' "
                "AND TABLE_NAME = 'members'".format(schema),
            )[0][0]
        ):
            return schema
        time.sleep(3)
    raise AssertionError("the repository of {} in {} did not build".format(FORM, code))


def _import(testapp, login, code, content, file_name="lists.json"):
    return testapp.post(
        "/user/{}/project/{}/caselists/import".format(login, code),
        status=302,
        upload_files=[("listfile", file_name, content)],
    )


def _refused(testapp, res, text):
    """The import was refused, and the lists page says why."""
    assert "FS_error" in res.headers
    page = res.follow()
    assert text.encode() in page.body, text


def _lists_of(config, project_id):
    return _query(
        config,
        "SELECT list_id, source_project, source_form, source_table, label_column, "
        "list_key_column, list_active FROM publishedlist "
        "WHERE project_id = '{}'".format(project_id),
    )


def _properties_of(config, project_id):
    return _query(
        config,
        "SELECT form_id, table_name, property_name, property_type, property_default "
        "FROM tableproperty WHERE project_id = '{}'".format(project_id),
    )


def t_e_s_t_list_transfer(test_object):
    login = test_object.randonLogin
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "list_transfer")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)
    xlsx = os.path.join(work, FORM + ".xlsx")
    _workbook(xlsx)

    # --- The example, in a project of its own ----------------------------------
    source_id = _new_project(testapp, config, login, SOURCE)
    _upload(testapp, login, SOURCE, xlsx)
    _build(testapp, config, login, SOURCE, source_id)
    lists = "/user/{}/project/{}/caselists".format(login, SOURCE)
    for list_data in (
        {
            "list_id": "farmers",
            "list_filename": "farmers.csv",
            "source_table": "maintable",
            "label_column": "farmer_name",
            "list_active": "1",
            "list_key_column": "",
        },
        {
            "list_id": "counties",
            "list_filename": "counties.csv",
            "source_table": "maintable",
            "label_column": "county",
            "list_active": "1",
            "list_key_column": "county",
        },
        {
            "list_id": "members",
            "list_filename": "members.csv",
            "source_table": "members",
            "label_column": "member_name",
            "list_active": "0",
            "list_key_column": "",
        },
    ):
        res = testapp.post(
            lists + "/add", dict(list_data, source_form=FORM), status=302
        )
        assert "FS_error" not in res.headers
    res = testapp.post(
        "/user/{}/project/{}/caseproperties?form={}&table=maintable".format(
            login, SOURCE, FORM
        ),
        {
            "add_property": "1",
            "property_name": "stage",
            "property_type": "string",
            "property_default": "registered",
            "property_desc": "Where the farmer is",
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    for list_id, column_name, column_as in (
        ("farmers", "county", "region"),
        ("farmers", "property:stage", ""),
        ("members", "parent_rowuuid", "farmer"),
    ):
        res = testapp.post(
            lists + "/{}/edit".format(list_id),
            {"add_column": "1", "column_name": column_name, "column_as": column_as},
            status=302,
        )
        assert "FS_error" not in res.headers
    rules = {
        "condition": "AND",
        "rules": [{"id": "p.stage", "operator": "equal", "value": "registered"}],
    }
    res = testapp.post(
        lists + "/farmers/edit",
        {"change_filter": "1", "filter_rules": json.dumps(rules)},
        status=302,
    )
    assert "FS_error" not in res.headers

    # --- Exported -----------------------------------------------------------------
    stage = {
        "form_id": FORM,
        "table": "maintable",
        "name": "stage",
        "type": "string",
        "default": "registered",
        "description": "Where the farmer is",
    }
    farmers = {
        "code": "farmers",
        "file_name": "farmers.csv",
        "form_id": FORM,
        "table": "maintable",
        "label_column": "farmer_name",
        "key_column": None,
        "rows": "active",
        "columns": [
            {"column": "county", "served_as": "region"},
            {"property": "stage"},
        ],
        "filter": rules,
    }
    res = testapp.get(lists + "/farmers/export?format=json", status=200)
    assert res.headers["Content-Disposition"] == 'attachment; filename="farmers.json"'
    one = json.loads(res.body)
    assert one["lists"] == [farmers], one
    assert one["properties"] == [stage], one

    res = testapp.get(lists + "/export?format=yaml", status=200)
    assert res.headers[
        "Content-Disposition"
    ] == 'attachment; filename="{}_lists.yaml"'.format(SOURCE)
    every = yaml.safe_load(res.body)
    by_code = {a["code"]: a for a in every["lists"]}
    assert sorted(by_code) == ["counties", "farmers", "members"], by_code
    assert by_code["farmers"] == farmers
    assert by_code["counties"]["key_column"] == "county"
    assert by_code["members"]["rows"] == "inactive"
    assert by_code["members"]["columns"] == [
        {"column": "parent_rowuuid", "served_as": "farmer"}
    ]
    # Two lists from maintable: its property once; members has none.
    assert every["properties"] == [stage], every["properties"]
    every_yaml = res.body

    # --- Imported into a project that is not ready, then is ----------------------
    target_id = _new_project(testapp, config, login, TARGET)
    res = _import(testapp, login, TARGET, every_yaml, "lists.yaml")
    _refused(testapp, res, "does not exist in this project")

    _upload(testapp, login, TARGET, xlsx)
    res = _import(testapp, login, TARGET, every_yaml, "lists.yaml")
    _refused(testapp, res, "has no repository yet")
    target_schema = _build(testapp, config, login, TARGET, target_id)

    broken = copy.deepcopy(one)
    broken["lists"][0]["columns"].append({"column": "nope"})
    res = _import(testapp, login, TARGET, json.dumps(broken).encode())
    _refused(testapp, res, "Nothing was imported")
    assert _lists_of(config, target_id) == []
    assert _properties_of(config, target_id) == []

    res = _import(testapp, login, TARGET, b"just text")
    _refused(testapp, res, "The file is not a list export of FormShare")
    res = testapp.post(
        "/user/{}/project/{}/caselists/import".format(login, TARGET), status=302
    )
    _refused(testapp, res, "Choose a JSON or YAML file to import")

    res = _import(testapp, login, TARGET, every_yaml, "lists.yaml")
    assert "FS_error" not in res.headers
    page = res.follow()
    assert b"Imported 3 list(s): " in page.body
    assert b"Created 1 propert(ies): stage." in page.body

    # The lists follow the target project's form, not the source's.
    assert sorted(tuple(r) for r in _lists_of(config, target_id)) == [
        ("counties", target_id, FORM, "maintable", "county", "county", 1),
        ("farmers", target_id, FORM, "maintable", "farmer_name", None, 1),
        ("members", target_id, FORM, "members", "member_name", None, 0),
    ]
    assert [tuple(r) for r in _properties_of(config, target_id)] == [
        (FORM, "maintable", "stage", "string", "registered")
    ]
    stored_rules, stored_sql = _query(
        config,
        "SELECT filter_rules, filter_sql FROM publishedlist "
        "WHERE project_id = '{}' AND list_id = 'farmers'".format(target_id),
    )[0]
    assert json.loads(stored_rules) == rules
    assert "HEX({p}`stage`)" in stored_sql, stored_sql
    assert _query(
        config,
        "SELECT COUNT(*) FROM information_schema.COLUMNS WHERE TABLE_SCHEMA = '{}' "
        "AND TABLE_NAME = 'maintable_properties' AND COLUMN_NAME = 'stage'".format(
            target_schema
        ),
    )[0][0], "the property should be a column of the target's repository"
    res = testapp.get(
        "/user/{}/project/{}/caselists/farmers/sample".format(login, TARGET),
        status=200,
    )
    assert res.body.decode("utf-8").splitlines()[0] == (
        '"name","label","region","stage"'
    ), res.body[:200]

    # What the target exports is what the source exported.
    res = testapp.get(
        "/user/{}/project/{}/caselists/export?format=json".format(login, TARGET),
        status=200,
    )
    exported = json.loads(res.body)
    assert exported["properties"] == every["properties"]
    assert sorted(exported["lists"], key=lambda a: a["code"]) == sorted(
        every["lists"], key=lambda a: a["code"]
    )

    # A second import finds every code in use, and writes nothing.
    res = _import(testapp, login, TARGET, every_yaml, "lists.yaml")
    _refused(testapp, res, "its code is in use")
    assert len(_lists_of(config, target_id)) == 3

    testapp.get(
        "/user/{}/project/{}/caselists/import".format(login, TARGET), status=404
    )
    testapp.get(
        "/user/{}/project/{}/caselists/ghost/export".format(login, TARGET),
        status=404,
    )
