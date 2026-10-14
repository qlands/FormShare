"""Field encryption's formats (processes/db/encryption.py), without a server.

docs/formshare_field_encryption/: an answer to a marked question is sealed to
the project's public key with libsodium's sealed box. What FormShare seals,
RSTools and the plugin must open, and what RSTools sealed -- the vectors in
resources/encryption/test_vectors.json, made through its real pipeline --
FormShare must recognise. The server never opens anything, so the opening
here is the test's own, with the test key the vectors publish.

The last tests run RSTools' jxformtomysql and xmltojson when they are
installed: FormShare seals a submission's XML before anything is written, so
it must choose the answers xmltojson would have chosen, and xmltojson must
keep what FormShare sealed.
"""

import json
import os
import subprocess

import nacl.bindings
import openpyxl
import pytest
from lxml import etree
from nacl.public import PrivateKey, SealedBox

from formshare.processes.db.encryption import (
    EncryptionFormatError,
    b64u_decode,
    b64u_encode,
    check_grant,
    check_key_signature,
    check_locked_key,
    check_unlock_params,
    format_member_key,
    format_public_key,
    format_sign_public_key,
    is_envelope,
    key_id,
    member_name,
    marked_columns,
    marked_xmlcodes,
    parse_grant,
    parse_member_key,
    parse_public_key,
    parse_sign_public_key,
    seal,
    seal_submission,
    seal_xml,
)
from formshare.tests.steps.encryption_keys import (
    _grant,
    _kdf,
    _key_taken,
    _sign_public_key,
    _signed,
)

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "resources", "encryption", "test_vectors.json")) as f:
    VECTORS = json.load(f)
PUBLIC_KEY = VECTORS["public_key"]
PRIVATE_KEY = b64u_decode(VECTORS["private_key_b64url"])
# Signed key versions and version 2 grants (README 3.6), made with libsodium's
# own key derivation from the plugin's vector keys: what the vault must make.
with open(os.path.join(HERE, "resources", "encryption", "signing_vectors.json")) as f:
    SIGNING = json.load(f)


def _open(envelope, private_key=PRIVATE_KEY):
    box = b64u_decode(envelope.split(":")[3])
    padded = SealedBox(PrivateKey(private_key)).decrypt(box)
    return nacl.bindings.sodium_unpad(padded, 32).decode("utf-8")


def _new_key():
    private = PrivateKey.generate()
    return bytes(private), format_public_key(bytes(private.public_key))


# --- The vectors -------------------------------------------------------------


def test_what_rstools_sealed_is_an_envelope_and_opens_to_the_answer():
    for vector in VECTORS["vectors"]:
        assert is_envelope(vector["envelope"]), vector["value"]
        assert _open(vector["envelope"]) == vector["value"]
        box = b64u_decode(vector["envelope"].split(":")[3])
        assert len(box) - 48 == vector["padded_length"]


def test_what_only_looks_like_an_envelope_is_not_one():
    for text in VECTORS["not_envelopes"]:
        assert not is_envelope(text), text
    envelope = VECTORS["vectors"][0]["envelope"]
    prefix, version, kid, box = envelope.split(":")
    for text in (
        None,
        42,
        "yes",
        ":".join([prefix, version, kid.upper(), box]),  # the id in capitals
        ":".join([prefix, version, kid[:15], box]),  # an id cut short
        ":".join([prefix, version, kid, box, "x"]),  # one part too many
        ":".join([prefix, version, kid, box[:-4]]),  # three bytes short
        ":".join([prefix, version, kid, box.replace("_", "/")]),  # base64
        ":".join([prefix, version, kid, b64u_encode(bytes(48))]),  # no block
        ":".join([prefix, version, kid, b64u_encode(bytes(48 + 31))]),
    ):
        assert not is_envelope(text), text
    assert is_envelope(":".join([prefix, version, kid, b64u_encode(bytes(80))]))


def test_what_formshare_seals_opens_to_the_answer_as_it_was_given():
    for vector in VECTORS["vectors"]:
        envelope = seal(vector["value"], PUBLIC_KEY)
        assert is_envelope(envelope)
        assert envelope.startswith("fsenc:1:" + VECTORS["key_id"] + ":")
        assert _open(envelope) == vector["value"]
        # Sealed to a random key each time, but padded alike: the same size
        # as what RSTools made of the same answer.
        assert envelope != vector["envelope"]
        assert len(envelope) == len(vector["envelope"])
    # Padding hides the length of a short answer.
    assert len(seal("yes", PUBLIC_KEY)) == len(seal("no", PUBLIC_KEY)) == 132


def test_an_empty_answer_and_an_envelope_are_left_as_they_are():
    assert seal(None, PUBLIC_KEY) is None
    assert seal("", PUBLIC_KEY) == ""
    assert seal("  \n", PUBLIC_KEY) == "  \n"
    envelope = VECTORS["vectors"][0]["envelope"]
    assert seal(envelope, PUBLIC_KEY) == envelope
    # A lookalike is an answer, and is sealed like any other.
    lookalike = "fsenc:1:0123456789abcdef:abc"
    assert _open(seal(lookalike, PUBLIC_KEY)) == lookalike
    with pytest.raises(TypeError):
        seal(34, PUBLIC_KEY)


# --- Keys, member keys, grants, locked keys -----------------------------------


def test_a_public_key_must_be_its_own():
    assert parse_public_key(PUBLIC_KEY) == (
        VECTORS["key_id"],
        b64u_decode(VECTORS["public_key_b64url"]),
    )
    assert parse_public_key("  " + PUBLIC_KEY + "\n")[0] == VECTORS["key_id"]
    kid = VECTORS["key_id"]
    key = VECTORS["public_key_b64url"]
    other = b64u_encode(bytes(32))
    for text in (
        None,
        "",
        "fsenc-mk:1:" + key,
        "fsenc-pk:2:{}:{}".format(kid, key),
        "fsenc-pk:1:{}:{}".format(kid, key[:-1]),
        "fsenc-pk:1:{}:{}=".format(kid, key),
        "fsenc-pk:1:{}:{}".format(kid, other),  # another key under this id
        "fsenc-pk:1:{}".format(key),
    ):
        with pytest.raises(EncryptionFormatError):
            parse_public_key(text)
    assert format_public_key(b64u_decode(key)) == PUBLIC_KEY
    assert key_id(b64u_decode(key)) == kid


def test_member_keys_grants_and_locked_keys_have_their_sizes():
    member = "fsenc-mk:1:" + b64u_encode(bytes(range(32)))
    assert parse_member_key(member) == bytes(range(32))
    for text in ("fsenc-mk:1:" + b64u_encode(bytes(31)), "fsenc-mk:2:x", PUBLIC_KEY):
        with pytest.raises(EncryptionFormatError):
            parse_member_key(text)

    kid = VECTORS["key_id"]
    sealed, signature = b64u_encode(bytes(80)), b64u_encode(bytes(64))
    grant = "fsenc-grant:3:{}:{}:{}".format(kid, sealed, signature)
    assert parse_grant(grant) == kid
    for text in (
        "fsenc-grant:3:{}:{}:{}".format(kid, b64u_encode(bytes(81)), signature),
        "fsenc-grant:3:{}:{}:{}".format(kid.upper(), sealed, signature),
        "fsenc-grant:3:{}:{}:{}".format(kid, sealed, b64u_encode(bytes(63))),
        "fsenc-grant:3:{}:{}".format(kid, sealed),
        # Version 1 had no signature: nothing shows who made it. Version 2's
        # named the key but not the member (formshare.md 15.1).
        "fsenc-grant:1:{}:{}".format(kid, sealed),
        "fsenc-grant:2:{}:{}:{}".format(kid, sealed, signature),
        VECTORS["vectors"][0]["envelope"],
    ):
        with pytest.raises(EncryptionFormatError):
            parse_grant(text)

    locked = {"nonce": b64u_encode(bytes(24)), "box": b64u_encode(bytes(48))}
    assert json.loads(check_locked_key(locked)) == locked
    assert json.loads(check_locked_key(json.dumps(locked))) == locked
    for bad in (
        {"nonce": b64u_encode(bytes(23)), "box": locked["box"]},
        {"nonce": locked["nonce"], "box": b64u_encode(bytes(64))},
        dict(locked, extra="x"),
        "not json",
    ):
        with pytest.raises(EncryptionFormatError):
            check_locked_key(bad)


def test_a_versions_signing_key_is_derived_as_libsodium_derives_it():
    # The vectors' seeds came from libsodium's crypto_kdf_derive_from_key; the
    # suite's browser derives them with BLAKE2b, as libsodium defines it.
    for version in (SIGNING["first_version"], SIGNING["second_version"]):
        private = b64u_decode(version["private_key_b64url"])
        assert b64u_encode(_kdf("fskeysig", private)) == version["seed_b64url"]
        assert _sign_public_key(PrivateKey(private)) == version["sign_public_key"]
        signed = _signed(PrivateKey(private), PrivateKey(private))
        assert signed["sign_public_key"] == version["sign_public_key"]
    first = SIGNING["first_version"]
    # The first version vouches for itself, as set-up signs it.
    assert _signed(PrivateKey(b64u_decode(first["private_key_b64url"]))) == {
        "sign_public_key": first["sign_public_key"],
        "key_signature": first["key_signature"],
    }


def test_a_key_version_is_vouched_for_by_the_one_it_replaced():
    first, second = SIGNING["first_version"], SIGNING["second_version"]
    check_key_signature(
        first["key_signature"],
        first["key_id"],
        first["sign_public_key"],
        first["public_key"],
    )
    check_key_signature(
        second["key_signature"],
        second["signed_by"],
        first["sign_public_key"],
        second["public_key"],
    )
    refused = (
        # The new version vouching for itself.
        (
            second["key_signature"],
            second["key_id"],
            second["sign_public_key"],
            second["public_key"],
        ),
        # The first version's signature, said to vouch for another key.
        (
            first["key_signature"],
            first["key_id"],
            first["sign_public_key"],
            second["public_key"],
        ),
        # A signing key that is not the signer's.
        (
            second["key_signature"],
            first["key_id"],
            second["sign_public_key"],
            second["public_key"],
        ),
        (
            b64u_encode(bytes(64)),
            first["key_id"],
            first["sign_public_key"],
            first["public_key"],
        ),
        (None, first["key_id"], first["sign_public_key"], first["public_key"]),
    )
    for signature, signer, sign_key, public_key in refused:
        with pytest.raises(EncryptionFormatError):
            check_key_signature(signature, signer, sign_key, public_key)


def test_a_grant_is_signed_by_its_version_for_its_member():
    first, second = SIGNING["first_version"], SIGNING["second_version"]
    grant = SIGNING["grant"]
    member, key = grant["member"], grant["member_public_key"]
    assert member == member_name("user", "vector_member")
    check_grant(grant["grant"], member, key, first["sign_public_key"])
    # A member key stored with spaces around it reads the same.
    check_grant(grant["grant"], member, " " + key + "\n", first["sign_public_key"])
    other = "fsenc-mk:1:" + b64u_encode(bytes(PrivateKey.generate().public_key))
    for name, member_key, sign_key in (
        (member, other, first["sign_public_key"]),  # another key
        # Another person, with this key put in beside the grant (15.1).
        (member_name("user", "someone_else"), key, first["sign_public_key"]),
        (member_name("recovery", "vector_member"), key, first["sign_public_key"]),
        (member, key, second["sign_public_key"]),  # another version's
        (None, key, first["sign_public_key"]),
    ):
        with pytest.raises(EncryptionFormatError):
            check_grant(grant["grant"], name, member_key, sign_key)
    the_key_id, sealed, _signature = grant["grant"].split(":")[2:]
    with pytest.raises(EncryptionFormatError):
        check_grant(
            "fsenc-grant:3:{}:{}:{}".format(the_key_id, sealed, b64u_encode(bytes(64))),
            member,
            key,
            first["sign_public_key"],
        )
    # The suite's browser makes grants core accepts, and signs nothing it
    # should not.
    private = PrivateKey(b64u_decode(first["private_key_b64url"]))
    check_grant(_grant(private, member, key), member, key, first["sign_public_key"])
    with pytest.raises(EncryptionFormatError):
        check_grant(
            _grant(private, member, key, PrivateKey.generate()),
            member,
            key,
            first["sign_public_key"],
        )
    assert format_member_key(parse_member_key(key)) == key


def test_a_member_is_named_by_who_they_are():
    assert member_name("user", "alice") == "user:alice"
    assert member_name("assistant", "4c1d") == "assistant:4c1d"
    assert member_name("recovery", "p1") == "recovery:p1"
    for member_type, identity in (("partner", "x"), ("user", ""), ("user", None)):
        with pytest.raises(EncryptionFormatError):
            member_name(member_type, identity)


def test_a_removal_recorded_before_keys_were_named_takes_none():
    # The plugin's flag by key, as the suite writes it (formshare.md 17.1). A
    # removal from before 17.1 has no member_key_ids (formshare.md 18.1).
    old_removal = {
        "event_type": "encryption_removed",
        "event_detail": {"key_ids": ["9653f4079c692b1e"], "grants": 3},
    }
    grant = {"event_type": "grant", "event_detail": {"member_key_id": "abcd"}}
    revoke = {"event_type": "revoke", "event_detail": {"member_key_id": "abcd"}}
    assert not _key_taken([old_removal, grant], "abcd")
    assert _key_taken([old_removal, revoke, grant], "abcd")
    named = dict(old_removal, event_detail={"member_key_ids": ["abcd"]})
    assert _key_taken([named, grant], "abcd")


def test_a_signing_key_names_its_version():
    first = SIGNING["first_version"]
    the_key_id, key = parse_sign_public_key(first["sign_public_key"])
    assert the_key_id == first["key_id"] and len(key) == 32
    assert format_sign_public_key(the_key_id, key) == first["sign_public_key"]
    encoded = first["sign_public_key"].split(":")[3]
    for text in (
        "fsenc-spk:2:{}:{}".format(the_key_id, encoded),
        "fsenc-spk:1:{}:{}".format(the_key_id.upper(), encoded),
        "fsenc-spk:1:{}:{}".format(the_key_id, b64u_encode(bytes(31))),
        "fsenc-pk:1:{}:{}".format(the_key_id, encoded),
        None,
    ):
        with pytest.raises(EncryptionFormatError):
            parse_sign_public_key(text)


def test_an_unlock_method_is_never_weaker_than_the_design_allows():
    salt = b64u_encode(bytes(16))
    good = {"salt": salt, "opslimit": 3, "memlimit": 64 * 1024 * 1024}
    good["alg"] = "argon2id13"
    assert check_unlock_params("passphrase", dict(good, other=1)) == good
    for weaker in (
        dict(good, opslimit=2),
        dict(good, memlimit=32 * 1024 * 1024),
        dict(good, alg="argon2i13"),
        dict(good, salt=b64u_encode(bytes(8))),
        dict(good, opslimit="3"),
    ):
        with pytest.raises(EncryptionFormatError):
            check_unlock_params("passphrase", weaker)

    prf = {"prf_salt": b64u_encode(bytes(32))}
    credential = b64u_encode(bytes(64))
    assert check_unlock_params("passkey", prf, credential) == prf
    for params, cred in (
        (prf, None),
        (prf, "not+base64"),
        ({"prf_salt": b64u_encode(bytes(16))}, credential),
    ):
        with pytest.raises(EncryptionFormatError):
            check_unlock_params("passkey", params, cred)
    with pytest.raises(EncryptionFormatError):
        check_unlock_params("password", good)


# --- Sealing a submission ------------------------------------------------------

MARKED = {"farmer_name", "hh/phone", "member/mem_name", "member/contact/mem_phone"}


def test_a_submissions_json_is_sealed_where_its_questions_are_marked():
    private, public = _new_key()
    lookalike = "fsenc:1:0123456789abcdef:abc"
    envelope = seal("from a case list", public)
    data = {
        "_xform_id_string": "enc",
        "i_d": "A1",
        "farmer_name": 'Ana María O\'Brien; "Jr"',
        "hh/phone": 722000000,
        "hh/notes": "kept in clear",
        "member": [
            {
                "member/mem_name": "Wanjiru",
                "member/mem_age": "12",
                "member/contact/mem_phone": lookalike,
            },
            {
                "member/mem_name": envelope,
                "member/mem_age": "9",
                "member/contact/mem_phone": "",
            },
            {"member/mem_name": None, "member/mem_age": "3"},
        ],
    }
    assert seal_submission(data, MARKED, public) == 4
    assert _open(data["farmer_name"], private) == 'Ana María O\'Brien; "Jr"'
    assert _open(data["hh/phone"], private) == "722000000"
    first, second, third = data["member"]
    assert _open(first["member/mem_name"], private) == "Wanjiru"
    assert _open(first["member/contact/mem_phone"], private) == lookalike
    assert second["member/mem_name"] == envelope
    assert second["member/contact/mem_phone"] == ""
    assert third["member/mem_name"] is None
    assert (data["i_d"], data["hh/notes"], first["member/mem_age"]) == (
        "A1",
        "kept in clear",
        "12",
    )
    # With nothing marked nothing is touched, and no key is needed.
    untouched = {"farmer_name": "Ana"}
    assert seal_submission(untouched, set(), None) == 0
    assert untouched == {"farmer_name": "Ana"}


SUBMISSION = """<?xml version="1.0"?>
<data id="enc" version="1" xmlns:orx="http://openrosa.org/xforms">
<i_d>A1</i_d>
<farmer_name>Ana María O'Brien; "Jr"</farmer_name>
<hh>
<phone>0722 000 000</phone>
<notes>kept in clear</notes>
</hh>
<member>
<mem_name>Wanjiru</mem_name>
<contact><mem_phone><!-- typed by hand -->07<![CDATA[ <11> ]]>11</mem_phone></contact>
</member>
<member>
<mem_name>   </mem_name>
<contact><mem_phone>{envelope}</mem_phone></contact>
</member>
<orx:meta><orx:instanceID>uuid:5e1d</orx:instanceID></orx:meta>
</data>
"""


def test_a_submissions_xml_is_sealed_as_xmltojson_reads_it():
    private, public = _new_key()
    envelope = seal("from a case list", public)
    received = SUBMISSION.format(envelope=envelope).encode("utf-8")
    sealed, count = seal_xml(received, MARKED, public)
    assert count == 4
    root = etree.fromstring(sealed)
    assert _open(root.findtext("farmer_name"), private) == 'Ana María O\'Brien; "Jr"'
    assert _open(root.findtext("hh/phone"), private) == "0722 000 000"
    first, second = root.findall("member")
    assert _open(first.findtext("mem_name"), private) == "Wanjiru"
    # All of its text, CDATA included, the comment not -- and the comment
    # is gone with what it annotated.
    phone = first.find("contact/mem_phone")
    assert _open(phone.text, private) == "07 <11> 11"
    assert len(phone) == 0
    # Blank, and sealed already: left as they are.
    assert second.findtext("mem_name") == "   "
    assert second.findtext("contact/mem_phone") == envelope
    # Nothing else changed.
    assert (root.findtext("i_d"), root.findtext("hh/notes")) == ("A1", "kept in clear")
    assert root.get("id") == "enc"
    assert b"<orx:instanceID>uuid:5e1d</orx:instanceID>" in sealed


def test_a_submission_with_nothing_to_seal_keeps_its_bytes():
    received = SUBMISSION.format(envelope="").encode("utf-8")
    _, public = _new_key()
    assert seal_xml(received, set(), None) == (received, 0)
    assert seal_xml(received, {"hh/notes_missing"}, public) == (received, 0)


def test_a_name_with_a_prefix_is_read_as_written():
    # xmltojson reads without namespaces: orx:meta/orx:instanceID is the key.
    private, public = _new_key()
    received = SUBMISSION.format(envelope="").encode("utf-8")
    sealed, count = seal_xml(received, {"orx:meta/orx:instanceID"}, public)
    assert count == 1
    found = etree.fromstring(sealed).find("{http://openrosa.org/xforms}meta")
    assert _open(found[0].text, private) == "uuid:5e1d"


def test_an_element_named_as_the_root_is_left_out_of_a_key():
    # xmltojson's quirk, kept: an ancestor named as the root is not in a key.
    private, public = _new_key()
    received = b"<data><data><x>1</x></data><y>2</y></data>"
    sealed, count = seal_xml(received, {"x", "y"}, public)
    root = etree.fromstring(sealed)
    assert count == 2
    assert _open(root.findtext("data/x"), private) == "1"


def test_the_marks_are_read_from_the_manifest_and_create_xml(tmp_path):
    (tmp_path / "manifest.xml").write_text(
        '<?xml version="1.0"?><XMLSchemaManifest><table mysqlcode="maintable" '
        'xmlcode="main"><field mysqlcode="farmer_name" xmlcode="farmer_name" '
        'fs_encrypted="true"/><field mysqlcode="i_d" xmlcode="i_d"/>'
        '<field mysqlcode="rowuuid" xmlcode="NONE" fs_encrypted="true"/>'
        '<table mysqlcode="member" xmlcode="member"><field mysqlcode="mem_name" '
        'xmlcode="member/mem_name" fs_encrypted="true"/></table></table>'
        "</XMLSchemaManifest>"
    )
    assert marked_xmlcodes(str(tmp_path / "manifest.xml")) == {
        "farmer_name",
        "member/mem_name",
    }
    (tmp_path / "create.xml").write_text(
        '<?xml version="1.0"?><XMLSchemaStructure><lkptables><table name="lkpx">'
        '<field name="x_cod"/></table></lkptables><tables><table name="maintable">'
        '<field name="farmer_name" fs_encrypted="true"/><field name="i_d"/>'
        '<table name="member"><field name="mem_name" fs_encrypted="true"/>'
        "</table></table></tables></XMLSchemaStructure>"
    )
    assert marked_columns(str(tmp_path / "create.xml")) == {
        ("maintable", "farmer_name"),
        ("member", "mem_name"),
    }


# --- Against the installed RSTools ---------------------------------------------


def _odktools_path():
    """Where the suite's configuration says RSTools is: ODKTOOLS_PATH, as CI
    sets it, or the generated test_config.json."""
    if os.environ.get("ODKTOOLS_PATH"):
        return os.environ["ODKTOOLS_PATH"]
    config = os.path.join(HERE, "test_config.json")
    if not os.path.exists(config):
        return ""
    with open(config) as config_file:
        return json.load(config_file).get("odktools.path", "")


def _contract_form(tmp_path):
    """A form with marked questions at the top, in a group, in a repeat and
    in a group inside it, as FormShare turns a workbook into what RSTools
    reads: the XForm, the survey JSON, then jxformtomysql's manifest."""
    from pyxform import xls2xform
    from pyxform.xls2json import parse_file_to_json

    tool = os.path.join(_odktools_path(), "JXFormToMysql", "jxformtomysql")
    if not os.path.exists(tool):
        pytest.skip("jxformtomysql is not installed")
    book = openpyxl.Workbook()
    survey = book.active
    survey.title = "survey"
    for row in (
        ["type", "name", "label", "fs_encrypted"],
        ["text", "i_d", "ID", ""],
        ["text", "farmer_name", "Name", "yes"],
        ["integer", "age", "Age", "Yes"],
        ["begin group", "hh", "Household", ""],
        ["text", "phone", "Phone", "yes"],
        ["text", "notes", "Notes", "no"],
        ["end group", "", "", ""],
        ["begin repeat", "member", "Member", ""],
        ["text", "mem_name", "Name", "yes"],
        ["integer", "mem_age", "Age", ""],
        ["begin group", "contact", "Contact", ""],
        ["text", "mem_phone", "Phone", "yes"],
        ["end group", "", "", ""],
        ["end repeat", "", "", ""],
    ):
        survey.append(row)
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Encryption contract", "enc", "1"])
    book.save(str(tmp_path / "enc.xlsx"))
    xls2xform.xls2xform_convert(str(tmp_path / "enc.xlsx"), str(tmp_path / "enc.xml"))
    survey_json = parse_file_to_json(str(tmp_path / "enc.xlsx"), warnings=[])
    (tmp_path / "enc.srv").write_text(json.dumps(survey_json))
    (tmp_path / "tmp").mkdir()
    run = subprocess.run(
        [
            tool,
            "-j",
            str(tmp_path / "enc.srv"),
            "-t",
            "maintable",
            "-v",
            "i_d",
            "-e",
            str(tmp_path / "tmp"),
            "-o",
            "m",
        ],
        cwd=str(tmp_path),
        capture_output=True,
    )
    assert run.returncode == 0, run.stdout
    manifest = etree.parse(str(tmp_path / "manifest.xml"))
    if not any(f.get("fs_encrypted") for f in manifest.iter("field")):
        pytest.skip("this jxformtomysql predates field encryption")
    return tmp_path / "enc.xml", tmp_path / "manifest.xml", tmp_path / "create.xml"


CONTRACT = """<?xml version="1.0"?>
<data id="enc" version="1"><i_d>A1</i_d><farmer_name>Ana María O'Brien; "Jr"</farmer_name><age>34</age><hh><phone>0722 000 000</phone><notes>kept in clear</notes></hh><member><mem_name>Wanjiru</mem_name><mem_age>12</mem_age><contact><mem_phone>0711 111 111</mem_phone></contact></member><member><mem_name>Otieno &amp; sons</mem_name><mem_age>9</mem_age><contact><mem_phone><![CDATA[07 <22> 22]]></mem_phone></contact></member><meta><instanceID>uuid:9f1c</instanceID></meta></data>
"""


def _xmltojson(xform, manifest, public, submission, out, sealed=None):
    tool = os.path.join(_odktools_path(), "XMLtoJSON", "xmltojson")
    args = [tool, "-i", str(submission), "-o", str(out), "-x", str(xform)]
    args += ["-m", str(manifest), "-k", public]
    if sealed is not None:
        args += ["-s", str(sealed)]
    run = subprocess.run(args, capture_output=True)
    assert run.returncode == 0, run.stdout
    with open(str(out)) as json_file:
        return json.load(json_file)


def _envelopes(data, prefix=""):
    """{(key, repeat position): envelope} of a submission's JSON."""
    found = {}
    for name, value in data.items():
        if isinstance(value, list):
            for position, item in enumerate(value):
                for (key, _), envelope in _envelopes(item).items():
                    found[(key, position)] = envelope
        elif is_envelope(value):
            found[(name, None)] = value
    return found


def _xmltojson_seals():
    tool = os.path.join(_odktools_path(), "XMLtoJSON", "xmltojson")
    if not os.path.exists(tool):
        return False
    run = subprocess.run([tool, "--help"], capture_output=True)
    return b"--publickey" in run.stdout


def test_formshare_seals_the_answers_xmltojson_would_seal(tmp_path):
    xform, manifest, create = _contract_form(tmp_path)
    if not _xmltojson_seals():
        pytest.skip("this xmltojson predates field encryption")
    marks = marked_xmlcodes(str(manifest))
    assert marks == {
        "farmer_name",
        "age",
        "hh/phone",
        "member/mem_name",
        "member/contact/mem_phone",
    }
    assert marked_columns(str(create)) == {
        ("maintable", "farmer_name"),
        ("maintable", "age"),
        ("maintable", "phone"),
        ("member", "mem_name"),
        ("member", "mem_phone"),
    }
    private, public = _new_key()
    (tmp_path / "plain.xml").write_text(CONTRACT)
    theirs = _envelopes(
        _xmltojson(xform, manifest, public, tmp_path / "plain.xml", tmp_path / "p.json")
    )
    sealed, count = seal_xml(CONTRACT.encode("utf-8"), marks, public)
    (tmp_path / "sealed.xml").write_bytes(sealed)
    ours = _envelopes(
        _xmltojson(
            xform,
            manifest,
            public,
            tmp_path / "sealed.xml",
            tmp_path / "s.json",
            tmp_path / "resealed.xml",
        )
    )
    # The same answers, sealed by either.
    assert count == len(theirs) == 7
    assert set(ours) == set(theirs)
    assert {k: _open(v, private) for k, v in ours.items()} == {
        k: _open(v, private) for k, v in theirs.items()
    }
    # xmltojson kept every envelope FormShare made: it found nothing in
    # clear to seal, in the JSON and in the XML it writes again.
    for (key, position), envelope in ours.items():
        if position is None:
            assert etree.fromstring(sealed).findtext(key) == envelope
    resealed = (tmp_path / "resealed.xml").read_bytes()
    for envelope in ours.values():
        assert envelope.encode() in resealed
    assert _open(ours[("member/contact/mem_phone", 1)], private) == "07 <22> 22"
    assert _open(ours[("member/mem_name", 1)], private) == "Otieno & sons"


# --- A form's marks through its life (formshare.md section 6) -------------------

from formshare.processes.db.encryption import (  # noqa: E402
    KEY_NEEDED,
    describe_invalid_encrypted_fields,
    encryption_change_message,
    field_is_encrypted,
    field_labelled_encrypted,
    key_needed_message,
    survey_marks,
)


def _same(text):
    return text


def test_the_marks_of_a_survey_are_read_as_jxformtomysql_reads_them(tmp_path):
    survey = {
        "type": "survey",
        "name": "data",
        "children": [
            {"type": "text", "name": "farmer", "fs_encrypted": "yes"},
            {"type": "integer", "name": "age", "FS_Encrypted": " TRUE "},
            {"type": "text", "name": "notes", "fs_encrypted": "no"},
            {"type": "text", "name": "plain"},
            {
                "type": "repeat",
                "name": "member",
                "fs_encrypted": "yes",
                "children": [
                    {"type": "text", "name": "mem_name", " fs_encrypted ": "Yes"},
                    {"type": "decimal", "name": "mem_weight", "fs_encrypted": True},
                ],
            },
        ],
    }
    (tmp_path / "form.srv").write_text(json.dumps(survey))
    # A repeat is not a question, and the mark is text, as pyxform writes it.
    assert survey_marks(str(tmp_path / "form.srv")) == ["farmer", "age", "mem_name"]


def test_what_a_dictionary_field_is():
    assert field_is_encrypted({"fs_encrypted": "true"}) == 1
    assert field_is_encrypted({"fs_encrypted": "yes"}) == 0
    assert field_is_encrypted({}) == 0
    # The retired label marks a field sensitive, never encrypted.
    assert field_labelled_encrypted({"formshare_encrypted": "yes"})
    assert field_labelled_encrypted({"formshare_encrypted": "1"})
    assert not field_labelled_encrypted({"formshare_encrypted": "no"})
    assert not field_labelled_encrypted({"fs_encrypted": "true"})


EXIT_39 = """<XMLResult>
  <XMLInvalidEncryptedFields total="3">
    <invalidField table="maintable" name="i_d" odktype="text" reason="key"/>
    <invalidField table="maintable" name="gender" odktype="select one" reason="select"/>
    <invalidField table="member" name="photo" odktype="image" reason="media"/>
  </XMLInvalidEncryptedFields>
</XMLResult>"""


def test_exit_39_is_told_question_by_question():
    lines = describe_invalid_encrypted_fields(etree.fromstring(EXIT_39), _same)
    assert lines == [
        'The question "i_d" cannot be encrypted: it is the form\'s key.',
        'The question "gender" cannot be encrypted: it is a select: its answer '
        "must match an option.",
        'The question "photo" of the table "member" cannot be encrypted: it is a '
        "file (photo, audio, video), which cannot be encrypted yet.",
    ]
    assert (
        describe_invalid_encrypted_fields(etree.fromstring("<XMLResult/>"), _same) == []
    )


def test_enc_says_which_way_the_mark_changed():
    unmarked = encryption_change_message(_same, "maintable", "farmer", True)
    rule = "A new version cannot change the encryption setting on a field."
    assert unmarked.startswith(rule)
    assert "encrypted in the repository and not in the new version" in unmarked
    assert "mark it fs_encrypted again" in unmarked
    marked = encryption_change_message(_same, "maintable", "notes", False)
    assert "not encrypted in the repository and is in the new version" in marked
    assert "remove fs_encrypted from it" in marked
    assert marked.startswith(rule)
    assert KEY_NEEDED not in range(0, 100)  # outside RSTools' codes
    assert "farmer, age" in key_needed_message(_same, ["farmer", "age"])


def test_a_case_variable_cannot_be_encrypted(tmp_path):
    from formshare.processes.odk.api import encrypted_case_variable

    (tmp_path / "create.xml").write_text(
        '<?xml version="1.0"?><XMLSchemaStructure><tables><table name="maintable">'
        '<field name="pid" type="varchar" key="true"/>'
        '<field name="hh_head" type="text" fs_encrypted="true"/>'
        '<field name="visit_date" type="date"/>'
        "</table></tables></XMLSchemaStructure>"
    )
    create = str(tmp_path / "create.xml")
    message = encrypted_case_variable(create, ["pid", "hh_head", "visit_date"], _same)
    assert message.startswith('The variable "hh_head" is encrypted')
    assert encrypted_case_variable(create, ["pid", None, "visit_date"], _same) is None


def test_a_filter_and_a_condition_say_what_they_read():
    from formshare.processes.db.actions import _encrypted_refusal, _rule_ids
    from formshare.processes.list_filter import referenced_columns

    rules = json.dumps(
        {
            "condition": "AND",
            "rules": [
                {"id": "t.farmer", "operator": "equal", "value": "Ann"},
                {
                    "condition": "OR",
                    "rules": [
                        {"id": "p.status", "operator": "equal", "value": "open"},
                        {"field": "t.age", "operator": "greater", "value": 3},
                    ],
                },
            ],
        }
    )
    assert referenced_columns(rules) == {"farmer", "age"}
    assert referenced_columns("not json") == set()

    class Catalogue:
        encrypted = {"v.hh_head", "k.phone"}
        encrypted_targets = {("case", "column", "phone")}

    def row(**changes):
        a_row = {
            "target_scope": "case",
            "target_kind": "property",
            "target_name": "status",
            "value_kind": "constant",
            "value": "open",
            "when_rules": None,
        }
        a_row.update(changes)
        return a_row

    assert _encrypted_refusal(row(), Catalogue, _same) is None
    written = _encrypted_refusal(
        row(target_kind="column", target_name="phone"), Catalogue, _same
    )
    assert (
        written.startswith('The column "phone" is encrypted') and "written" in written
    )
    copied = _encrypted_refusal(
        row(value_kind="variable", value="hh_head"), Catalogue, _same
    )
    assert "copied by an action" in copied
    condition = json.dumps(
        {"condition": "AND", "rules": [{"id": "k.phone", "operator": "is_null"}]}
    )
    assert _rule_ids(condition) == {"k.phone"}
    read = _encrypted_refusal(row(when_rules=condition), Catalogue, _same)
    assert read.startswith('The column "phone" is encrypted') and "condition" in read


def test_the_installed_jxformtomysql_refuses_a_mark_that_cannot_work(tmp_path):
    from pyxform.xls2json import parse_file_to_json

    tool = os.path.join(_odktools_path(), "JXFormToMysql", "jxformtomysql")
    if not os.path.exists(tool):
        pytest.skip("jxformtomysql is not installed")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "survey"
    for a_row in (
        ["type", "name", "label", "fs_encrypted"],
        ["text", "i_d", "ID", ""],
        ["select_one gender", "gender", "Gender", "yes"],
        ["geopoint", "home", "Home", "yes"],
    ):
        sheet.append(a_row)
    choices = book.create_sheet("choices")
    for a_row in (["list_name", "name", "label"], ["gender", "1", "Female"]):
        choices.append(a_row)
    settings = book.create_sheet("settings")
    settings.append(["form_id", "version"])
    settings.append(["refused", "1"])
    book.save(str(tmp_path / "refused.xlsx"))
    survey = parse_file_to_json(str(tmp_path / "refused.xlsx"), warnings=[])
    (tmp_path / "refused.srv").write_text(json.dumps(survey))
    (tmp_path / "tmp").mkdir()
    run = subprocess.run(
        [
            tool,
            "-j",
            str(tmp_path / "refused.srv"),
            "-t",
            "maintable",
            "-v",
            "i_d",
            "-e",
            str(tmp_path / "tmp"),
            "-o",
            "m",
            "-K",
        ],
        cwd=str(tmp_path),
        capture_output=True,
    )
    if run.returncode != 39:
        pytest.skip("this jxformtomysql predates field encryption")
    from formshare.processes.odk.api import machine_xml

    lines = describe_invalid_encrypted_fields(
        etree.fromstring(machine_xml(run.stdout)), _same
    )
    assert lines == [
        'The question "gender" cannot be encrypted: it is a select: its answer '
        "must match an option.",
        'The question "home" cannot be encrypted: it is a location.',
    ]


# --- The submission pipeline (formshare.md section 5) ----------------------------

from formshare.processes.db.encryption import (  # noqa: E402
    manifest_marks,
    no_key_message,
    seal_json_text,
    submission_digest,
)


def test_a_resend_is_recognised_without_its_answers():
    private, public = _new_key()
    sent = {
        "_xform_id_string": "enc",
        "pid": "A1",
        "farmer_name": "Ana",
        "member": [{"member/mem_name": "Wanjiru", "member/mem_age": "12"}],
        "meta/instanceID": "uuid:1",
    }
    first, second = json.loads(json.dumps(sent)), json.loads(json.dumps(sent))
    seal_submission(first, MARKED, public)
    seal_submission(second, MARKED, public)
    assert first["farmer_name"] != second["farmer_name"]  # sealing is random
    assert submission_digest(first) == submission_digest(second)
    # Anything in clear still tells two submissions apart.
    third = json.loads(json.dumps(sent))
    third["pid"] = "A2"
    seal_submission(third, MARKED, public)
    assert submission_digest(third) != submission_digest(first)
    assert submission_digest({"a": 1, "b": 2}) == submission_digest({"b": 2, "a": 1})


def test_a_pushed_json_is_kept_as_sent_unless_something_is_sealed():
    private, public = _new_key()
    raw = b'{"_xform_id_string": "enc",  "farmer_name": "Ana"}'
    assert seal_json_text(raw, {"notes"}, public) == (raw, 0)
    kept, sealed = seal_json_text(raw, {"farmer_name"}, public)
    assert sealed == 1
    assert _open(json.loads(kept)["farmer_name"], private) == "Ana"
    # A client that sealed for itself keeps its bytes.
    already = json.dumps({"farmer_name": seal("Ana", public)}).encode()
    assert seal_json_text(already, {"farmer_name"}, public) == (already, 0)


def test_a_forms_marks_come_from_its_manifest(tmp_path):
    repository = tmp_path / "forms" / "dir1" / "repository"
    repository.mkdir(parents=True)
    (repository / "manifest.xml").write_text(
        '<?xml version="1.0"?><XMLSchemaManifest><table mysqlcode="maintable">'
        '<field mysqlcode="farmer_name" xmlcode="farmer_name" fs_encrypted="true"/>'
        "</table></XMLSchemaManifest>"
    )
    assert manifest_marks(str(tmp_path), "dir1") == {"farmer_name"}
    assert manifest_marks(str(tmp_path), "nowhere") == set()
    assert "encryption is not set up" in no_key_message(_same)


def test_the_stored_command_hides_the_password():
    from formshare.processes.logging.loggerclass import masked_command

    assert masked_command(
        ["jsontomysql", "-u root", "-p s3cret", "-p=other", "-s FS_x"]
    ) == ("jsontomysql -u root -p xxxxxxxx -p=xxxxxxxx -s FS_x")
    # The password as an argument of its own, after "-p".
    assert masked_command(["jsontomysql", "-u", "root", "-p", "s3cret", "-s", "x"]) == (
        "jsontomysql -u root -p xxxxxxxx -s x"
    )


def test_a_plugin_sees_the_submission_sealed():
    from formshare.processes.odk.api import post_with_sealed_xml

    class Part:
        def __init__(self, filename, data):
            self.filename = filename
            self.file = data

    post = {
        "xml_submission_file": Part("submission.xml", b"<data>clear</data>"),
        "image": Part("photo.jpg", b"jpeg"),
        "*isIncomplete*": "yes",
    }
    seen = post_with_sealed_xml(post, "xml_submission_file", b"<data>sealed</data>")
    assert seen["xml_submission_file"].file.read() == b"<data>sealed</data>"
    assert seen["xml_submission_file"].filename == "submission.xml"
    assert seen["image"] is post["image"]
    assert seen["*isIncomplete*"] == "yes"


def test_an_import_is_unpacked_sealed(tmp_path):
    import io
    import zipfile

    from formshare.processes.odk.api import extract_sealed

    private, public = _new_key()
    archive_bytes = io.BytesIO()
    with zipfile.ZipFile(archive_bytes, "w") as archive:
        archive.writestr("one.json", json.dumps({"farmer_name": "Ana", "pid": "1"}))
        archive.writestr(
            "two/two.xml",
            '<?xml version="1.0"?><data id="enc"><farmer_name>Otieno</farmer_name>'
            "<pid>2</pid></data>",
        )
        archive.writestr("two/photo.jpg", b"jpeg")
        archive.writestr("../outside.json", json.dumps({"farmer_name": "Bad"}))
    target = tmp_path / "unpacked"
    target.mkdir()
    with zipfile.ZipFile(io.BytesIO(archive_bytes.getvalue())) as archive:
        extract_sealed(archive, str(target), {"farmer_name"}, public)
    one = json.loads((target / "one.json").read_text())
    assert _open(one["farmer_name"], private) == "Ana" and one["pid"] == "1"
    two = etree.fromstring((target / "two" / "two.xml").read_bytes())
    assert _open(two.findtext("farmer_name"), private) == "Otieno"
    assert (target / "two" / "photo.jpg").read_bytes() == b"jpeg"
    assert not (tmp_path / "outside.json").exists()
    # Nothing unpacked holds an answer in clear.
    for a_file in target.rglob("*"):
        if a_file.is_file():
            assert b"Ana" not in a_file.read_bytes()
            assert b"Otieno" not in a_file.read_bytes()


def test_a_fixed_submission_is_sealed_before_it_is_kept(tmp_path, monkeypatch):
    """A person fixing a submission may type an answer the form encrypts: the
    check-in seals it before it replaces the stored JSON or goes into the
    diff (formshare.md 5.4)."""
    import io

    from formshare.processes.odk import api

    private, public = _new_key()
    form_directory = tmp_path / "odk" / "forms" / "dir1"
    (form_directory / "repository").mkdir(parents=True)
    (form_directory / "repository" / "manifest.xml").write_text(
        '<?xml version="1.0"?><XMLSchemaManifest><table mysqlcode="maintable">'
        '<field mysqlcode="farmer_name" xmlcode="farmer_name" fs_encrypted="true"/>'
        "</table></XMLSchemaManifest>"
    )
    submissions = form_directory / "submissions"
    (submissions / "S1").mkdir(parents=True)
    (submissions / "S1.json").write_text(
        json.dumps({"pid": "S1", "farmer_name": seal("Ann", public)}, indent=4)
    )
    monkeypatch.setattr(api, "get_form_directory", lambda r, p, f: "dir1")
    monkeypatch.setattr(api, "get_active_project_key", lambda r, p: public)
    monkeypatch.setattr(api, "update_json_status", lambda *args: None)
    monkeypatch.setattr(api, "add_json_history", lambda *args: (True, ""))

    class Registry:
        settings = {"repository.path": str(tmp_path)}

    class Request:
        registry = Registry()

        @staticmethod
        def translate(text):
            return text

    fixed = io.BytesIO(json.dumps({"pid": "S1", "farmer_name": "Ann Fixed"}).encode())
    assert api.store_new_version(
        Request(), "u", "p", "f", "S1", "a", "seq1", fixed, "notes"
    ) == (0, "")
    stored = json.loads((submissions / "S1.json").read_text())
    assert _open(stored["farmer_name"], private) == "Ann Fixed"
    diff = (submissions / "S1" / "diffs" / "seq1" / "seq1.diff").read_text()
    assert "Ann Fixed" not in diff and "fsenc:1:" in diff

    # No key, nothing replaced.
    monkeypatch.setattr(api, "get_active_project_key", lambda r, p: None)
    before = (submissions / "S1.json").read_text()
    again = io.BytesIO(json.dumps({"pid": "S1", "farmer_name": "Clear"}).encode())
    res, message = api.store_new_version(
        Request(), "u", "p", "f", "S1", "a", "seq2", again, "notes"
    )
    assert res == 1 and "encryption is not set up" in message
    assert (submissions / "S1.json").read_text() == before


# --- Writes and errors (formshare.md section 7) -----------------------------------

from formshare.processes.db.encryption import (  # noqa: E402
    encrypted_edit_message,
    is_encrypted_refusal,
)


class _MySQLError(Exception):
    def __init__(self, errno, message):
        super().__init__(errno, message)
        self.errno = errno


class _Wrapped(Exception):
    """As SQLAlchemy wraps a driver's error: the driver's in .orig."""

    def __init__(self, original):
        super().__init__(str(original))
        self.orig = original


def test_an_answer_in_clear_in_an_encrypted_column_is_told_apart():
    refused = _MySQLError(
        3819, "Check constraint 'fsenc_0123456789abcdef' is violated."
    )
    assert is_encrypted_refusal(refused)
    assert is_encrypted_refusal(_Wrapped(refused))
    # Another CHECK, or another error, is not encryption's.
    assert not is_encrypted_refusal(
        _MySQLError(3819, "Check constraint 'other' is violated.")
    )
    assert not is_encrypted_refusal(
        _Wrapped(_MySQLError(1062, "Duplicate entry 'fsenc_'"))
    )
    assert not is_encrypted_refusal(ValueError("fsenc_"))
    assert "unlock it to edit" in encrypted_edit_message(_same)
    # The API leaves encrypted answers alone (/api_update ignores them).
    assert "does not change encrypted" in encrypted_edit_message(_same, for_api=True)


def test_a_name_from_a_request_is_quoted_only_when_it_is_one():
    from formshare.processes.submission.api import _quoted

    assert _quoted("maintable") == "`maintable`"
    assert _quoted("FS_8e3a_1") == "`FS_8e3a_1`"
    for bad in ("notes`; DROP TABLE x; --", "a b", "", None, "naïve"):
        with pytest.raises(ValueError):
            _quoted(bad)


def test_an_action_that_writes_an_encrypted_column_says_which():
    from formshare.processes.actions.host import ModuleError
    from formshare.processes.actions.server import _apply

    class Session:
        def __init__(self):
            self.committed = False

        def execute(self, statement, params=None):
            if "UPDATE" in str(statement) and "secret" in str(statement):
                raise _Wrapped(
                    _MySQLError(3819, "Check constraint 'fsenc_abc' is violated.")
                )

        def commit(self):
            self.committed = True

    class Request:
        dbsession = Session()

    plan = [
        {"table": "maintable", "column": "visits", "value": "2", "rowuuid": "r1"},
        {"table": "maintable", "column": "secret", "value": "clear", "rowuuid": "r1"},
    ]
    with pytest.raises(ModuleError) as refused:
        _apply(Request(), "FS_x", plan, "assistant")
    assert "maintable.secret is encrypted" in str(refused.value)
    assert not Request.dbsession.committed
