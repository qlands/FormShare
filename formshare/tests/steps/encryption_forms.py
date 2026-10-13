"""A form's encrypted questions through its life (formshare.md section 6).

A question marked fs_encrypted holds its answers sealed. What the form may
mark is RSTools' to decide, and what is done with a marked column FormShare's:

- a mark that cannot work (a select, a location) is refused at upload with
  jxformtomysql's exit 39, question by question;
- a form with marks gets no repository in a project without a key, and gets
  one once the project has a key; its dictionary then says which columns
  hold envelopes (field_encrypted, from create.xml), and the retired
  formshare_encrypted label makes a field sensitive and nothing more;
- a column that holds envelopes cannot be a published list's label, key or
  filter, though it may be served;
- a new version that marks a column the repository holds in clear, or
  unmarks a sealed one, is refused with mergeversions' ENC, which says which
  way the mark changed;
- a new version with marks cannot be merged in a project without a key.

The project's key is set up by core's functions, playing the browser as the
encryption_keys step does.
"""

import json
import os
import shutil
import time
import uuid

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import NullPool

from formshare.processes.db.encryption import (
    member_name,
    enroll_member,
    get_member,
    parse_public_key,
    set_up_project_encryption,
)
from .encryption_keys import _Person, _Request, _grant, _project_key, _signed
from .sql import get_form_details

CODE = "enc_forms"
NO_KEY = "enc_nokey"
HEADER = ["type", "name", "label", "fs_encrypted", "formshare_encrypted"]
MARKED_FORM = [
    ["text", "pid", "ID", "", ""],
    ["text", "farmer_name", "Farmer", "yes", ""],
    ["integer", "age", "Age", "Yes", ""],
    ["text", "phone", "Phone", "", "yes"],
    ["text", "notes", "Notes", "", ""],
    ["begin repeat", "member", "Member", "", ""],
    ["text", "mem_name", "Name", "yes", ""],
    ["end repeat", "", "", "", ""],
]


def _workbook(path, form_id, rows, choices=None):
    book = openpyxl.Workbook()
    survey = book.active
    survey.title = "survey"
    survey.append(HEADER)
    for a_row in rows:
        survey.append(a_row)
    if choices:
        sheet = book.create_sheet("choices")
        for a_row in choices:
            sheet.append(a_row)
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append([form_id, form_id, "1"])
    book.save(path)
    return path


def _query(engine, sql, **params):
    with engine.connect() as connection:
        return connection.execute(sql, params).fetchall()


def _new_project(testapp, login, code):
    res = testapp.post(
        "/user/{}/projects/add".format(login),
        {
            "project_code": code,
            "project_name": code,
            "project_abstract": "",
            "project_icon": "🔒",
            "project_hexcolor": "#9bbb59",
            "project_formlist_auth": 1,
        },
        status=302,
    )
    assert "FS_error" not in res.headers


def _project_id(engine, login, code):
    return _query(
        engine,
        "SELECT p.project_id FROM project p JOIN userproject u ON "
        "u.project_id = p.project_id WHERE p.project_code = %(code)s AND "
        "u.user_id = %(user)s AND u.access_type = 1",
        code=code,
        user=login,
    )[0][0]


def _assistant(testapp, login, code):
    coll_uuid = str(uuid.uuid4())
    coll_id = "enc_" + coll_uuid[-8:]
    res = testapp.post(
        "/user/{}/project/{}/assistants/add".format(login, code),
        {
            "coll_uuid": coll_uuid,
            "coll_id": coll_id,
            "coll_name": coll_id,
            "coll_password": "123",
            "coll_password2": "123",
            "coll_prjshare": 0,
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    return coll_id, coll_uuid


def _upload(testapp, login, code, path, form_pkey="pid"):
    return testapp.post(
        "/user/{}/project/{}/forms/add".format(login, code),
        {"form_pkey": form_pkey},
        status=302,
        upload_files=[("xlsx", path)],
    )


def _give(testapp, login, code, form, project_id, assistant):
    coll_id, coll_uuid = assistant
    res = testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(login, code, form),
        {
            "coll_id": "{}|{}|{}".format(project_id, coll_id, coll_uuid),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )
    assert "FS_error" not in res.headers


def _built(engine, config, project_id, form):
    """Waits for a form's repository; returns its schema."""
    for _ in range(100):
        schema = get_form_details(config, project_id, form)["form_schema"]
        if (
            schema
            and _query(
                engine,
                "SELECT COUNT(*) FROM information_schema.TABLES WHERE "
                "TABLE_SCHEMA = %(schema)s AND TABLE_NAME = 'maintable'",
                schema=schema,
            )[0][0]
        ):
            return schema
        time.sleep(3)
    raise AssertionError("the repository of {} did not build".format(form))


def _new_version(testapp, login, code, form, project_id, path):
    res = testapp.post(
        "/user/{}/project/{}/form/{}/merge".format(login, code, form),
        {"for_merging": "", "parent_project": project_id, "parent_form": form},
        status=302,
        upload_files=[("xlsx", path)],
    )
    assert "FS_error" not in res.headers


def t_e_s_t_encryption_forms(test_object):
    login = test_object.randonLogin
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "encryption_forms")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    session = sessionmaker(bind=engine)()
    try:
        _run(test_object, login, testapp, config, work, engine, session)
    finally:
        session.close()
        engine.dispose()


def _run(test_object, login, testapp, config, work, engine, session):
    _new_project(testapp, login, CODE)
    project_id = _project_id(engine, login, CODE)
    assistant = _assistant(testapp, login, CODE)
    form_url = "/user/{}/project/{}/form/{}".format(login, CODE, "enc_form")

    # --- A mark that cannot work: refused at upload (exit 39) ----------------
    refused = _workbook(
        os.path.join(work, "enc_refused.xlsx"),
        "enc_refused",
        [
            ["text", "pid", "ID", "", ""],
            ["select_one yesno", "agrees", "Agrees", "yes", ""],
            ["geopoint", "home", "Home", "yes", ""],
        ],
        [["list_name", "name", "label"], ["yesno", "1", "Yes"], ["yesno", "0", "No"]],
    )
    res = _upload(testapp, login, CODE, refused)
    assert "FS_error" in res.headers
    res = res.follow(status=200)
    assert b"cannot be encrypted" in res.body
    assert b"it is a select" in res.body and b"it is a location" in res.body

    # --- No repository without a key ------------------------------------------
    marked = _workbook(os.path.join(work, "enc_form.xlsx"), "enc_form", MARKED_FORM)
    res = _upload(testapp, login, CODE, marked)
    assert "FS_error" not in res.headers
    _give(testapp, login, CODE, "enc_form", project_id, assistant)
    res = testapp.post(
        form_url + "/repository/create",
        {"form_pkey": "pid", "start_stage1": ""},
        status=200,
    )
    assert b"Encryption is not set up for this project" in res.body
    assert b"farmer_name, age, mem_name" in res.body
    assert get_form_details(config, project_id, "enc_form")["form_schema"] is None

    # --- The owner sets up the project's key ---------------------------------
    request = _Request(session)
    owner, recovery = _Person(), _Person()
    # The owner may hold a key already (the encryption_keys step enrolled
    # them): the server checks a grant's form and key id, not whom it opens
    # for, so the step seals to whichever key the owner has.
    existing = get_member(request, "user", login)
    if existing is None:
        ok, owner_id = enroll_member(
            request, "user", login, owner.public_key, owner.unlock()
        )
        assert ok, owner_id
        owner_public = owner.public_key
    else:
        owner_id, owner_public = existing["member_id"], existing["public_key"]
    project_private, public_key = _project_key()
    ok, key_id = set_up_project_encryption(
        request,
        project_id,
        public_key,
        owner_id,
        _grant(project_private, member_name("user", login), owner_public),
        recovery.public_key,
        _grant(
            project_private,
            member_name("recovery", project_id),
            recovery.public_key,
        ),
        login,
        **_signed(project_private),
    )
    assert ok and key_id == parse_public_key(public_key)[0], key_id

    # --- Now the repository builds ---------------------------------------------
    res = testapp.post(
        form_url + "/repository/create",
        {"form_pkey": "pid", "start_stage1": ""},
        status=302,
    )
    assert "FS_error" not in res.headers
    schema = _built(engine, config, project_id, "enc_form")
    flags = {
        (table, field): (encrypted, sensitive)
        for table, field, encrypted, sensitive in _query(
            engine,
            "SELECT table_name, field_name, field_encrypted, field_sensitive FROM "
            "dictfield WHERE project_id = %(p)s AND form_id = 'enc_form'",
            p=project_id,
        )
    }
    assert flags[("maintable", "farmer_name")] == (1, 0)
    assert flags[("maintable", "age")] == (1, 0)
    assert flags[("member", "mem_name")] == (1, 0)
    assert flags[("maintable", "notes")] == (0, 0)
    # The retired label: sensitive, never encrypted.
    assert flags[("maintable", "phone")] == (0, 1)
    checks = _query(
        engine,
        "SELECT COUNT(*) FROM information_schema.CHECK_CONSTRAINTS WHERE "
        "CONSTRAINT_SCHEMA = %(schema)s AND CONSTRAINT_NAME LIKE 'fsenc_%%'",
        schema=schema,
    )[0][0]
    assert checks == 3, checks

    # --- What a published list may and may not do with a sealed column -------
    lists = "/user/{}/project/{}/caselists".format(login, CODE)
    new_list = {
        "list_id": "enc_farms",
        "list_filename": "enc_farms.csv",
        "source_form": "enc_form",
        "source_table": "maintable",
        "label_column": "farmer_name",
        "list_active": "1",
        "list_key_column": "",
    }
    res = testapp.post(lists + "/add", new_list, status=200)
    assert "FS_error" in res.headers
    assert b"farmer_name" in res.body and b"is encrypted" in res.body
    res = testapp.post(
        lists + "/add", dict(new_list, list_key_column="age"), status=200
    )
    assert "FS_error" in res.headers
    res = testapp.post(lists + "/add", dict(new_list, label_column="notes"), status=302)
    assert "FS_error" not in res.headers
    edit = lists + "/enc_farms/edit"
    for change in (
        {"change_label": "", "label_column": "farmer_name"},
        {"change_key": "", "list_key_column": "age"},
        {
            "change_filter": "",
            "filter_rules": json.dumps(
                {
                    "condition": "AND",
                    "rules": [
                        {
                            "id": "t.farmer_name",
                            "field": "t.farmer_name",
                            "operator": "equal",
                            "value": "Ann",
                        }
                    ],
                }
            ),
        },
    ):
        res = testapp.post(edit, change, status=200)
        assert "FS_error" in res.headers, change
        assert b"is encrypted" in res.body, change
    # A sealed column may be served: it goes out as its envelope.
    res = testapp.post(
        edit, {"add_column": "", "column_name": "farmer_name"}, status=302
    )
    assert "FS_error" not in res.headers

    # --- A new version may not change a mark (ENC) -----------------------------
    changed = [list(a_row) for a_row in MARKED_FORM]
    changed[1][3] = ""  # farmer_name: no longer encrypted
    changed[4][3] = "yes"  # notes: now encrypted
    _new_version(
        testapp,
        login,
        CODE,
        "enc_form",
        project_id,
        _workbook(os.path.join(work, "enc_form_b.xlsx"), "enc_form_b", changed),
    )
    res = testapp.get(
        "/user/{}/project/{}/form/{}".format(login, CODE, "enc_form_b"), status=200
    )
    assert b"Unable to merge" in res.body
    body = res.body.decode("utf-8")
    assert "farmer_name" in body and "is encrypted in the repository and not" in body
    assert "notes" in body and "is not encrypted in the repository and is" in body

    # --- Nor bring marks into a project without a key ----------------------------
    _new_project(testapp, login, NO_KEY)
    plain_project = _project_id(engine, login, NO_KEY)
    plain_rows = [[a_row[0], a_row[1], a_row[2], "", ""] for a_row in MARKED_FORM]
    _upload(
        testapp,
        login,
        NO_KEY,
        _workbook(os.path.join(work, "enc_plain.xlsx"), "enc_plain", plain_rows),
    )
    _give(
        testapp,
        login,
        NO_KEY,
        "enc_plain",
        plain_project,
        _assistant(testapp, login, NO_KEY),
    )
    res = testapp.post(
        "/user/{}/project/{}/form/{}/repository/create".format(
            login, NO_KEY, "enc_plain"
        ),
        {"form_pkey": "pid", "start_stage1": ""},
        status=302,
    )
    assert "FS_error" not in res.headers
    _built(engine, config, plain_project, "enc_plain")
    _new_version(
        testapp,
        login,
        NO_KEY,
        "enc_plain",
        plain_project,
        _workbook(
            os.path.join(work, "enc_plain_b.xlsx"),
            "enc_plain_b",
            plain_rows + [["text", "secret", "Secret", "yes", ""]],
        ),
    )
    res = testapp.get(
        "/user/{}/project/{}/form/{}".format(login, NO_KEY, "enc_plain_b"),
        status=200,
    )
    assert b"Unable to merge" in res.body
    assert b"This form encrypts the answers to secret" in res.body
