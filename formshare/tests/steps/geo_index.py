"""Traces and shapes in the dataset index, which the map reads.

A submission located by a geoshape or a geotrace had nothing in the dataset
index for the map to draw. FormShare now gives each document the first trace
or shape of the main table as GeoJSON (_geoshape), a point on it when no
geopoint gave one (_geolocation), and _geotype (processes/geometry.py):

- a submission in testing and one loaded into the repository are indexed
  alike;
- a boundary that crosses itself is repaired into a MultiPolygon that
  Elasticsearch takes, counted by a point inside it;
- a geopoint still places a document when the form has one;
- a shape query over a box finds the documents whose shapes cross it, as the
  map's will;
- index_geometries gives a document indexed without its shape that shape;
- a form located by a trace alone is geo-referenced: its page has the map,
  and the core map places the submission by the point on its trace.
"""

import os
import shutil
import time
import uuid

import openpyxl
from shapely.geometry import Point, shape
from sqlalchemy import create_engine
from sqlalchemy.pool import NullPool

from formshare.processes.elasticsearch.repository_index import (
    create_connection,
    get_index_name,
)
from formshare.scripts.indexgeometries import main as index_geometries
from .sql import get_form_details

FORM = "geo_index"
ROUTE_FORM = "geo_route"
BOWTIE = (
    "-1.000 36.000 0 0;-1.000 36.002 0 0;-1.002 36.000 0 0;"
    "-1.002 36.002 0 0;-1.000 36.000 0 0"
)
SQUARE = "-2 37 0 0;-2 37.01 0 0;-2.01 37.01 0 0;-2.01 37 0 0;-2 37 0 0"
ROUTE = "-3 38 0 0;-3.01 38.01 0 0;-3.02 38.03 0 0"
HOME = "-2.005 37.005 0 0"


def _workbook(path):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "i_d", "ID"])
    survey.append(["geopoint", "home", "Home"])
    survey.append(["geoshape", "plot", "Plot"])
    survey.append(["geotrace", "route", "Route"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Geometry in the index", FORM, "1"])
    wb.save(path)


def _submission(path, key, home="", plot="", route=""):
    xml = (
        "<?xml version='1.0'?>"
        '<data id="{form}" version="1">'
        "<i_d>{key}</i_d><home>{home}</home><plot>{plot}</plot>"
        "<route>{route}</route>"
        "<meta><instanceID>uuid:{uuid}</instanceID></meta></data>"
    ).format(form=FORM, key=key, home=home, plot=plot, route=route, uuid=uuid.uuid4())
    with open(path, "w", encoding="utf-8") as a_file:
        a_file.write(xml)


def _route_workbook(path):
    """A form whose only location is a geotrace."""
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "i_d", "ID"])
    survey.append(["geotrace", "route", "Route"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["A route alone", ROUTE_FORM, "1"])
    wb.save(path)


def _query(config, sql):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        return engine.execute(sql).fetchall()
    finally:
        engine.dispose()


def _inside(document):
    point = Point(document["_geolocation"]["lon"], document["_geolocation"]["lat"])
    return shape(document["_geoshape"]).buffer(1e-9).contains(point)


def t_e_s_t_geo_index(test_object):
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "geo_index")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)
    connection = create_connection(config)
    index_name = get_index_name(config)

    def push(key, **answers):
        xml = os.path.join(work, key + ".xml")
        _submission(xml, key, **answers)
        res = testapp.post(
            "/user/{}/project/{}/push".format(login, project),
            status=201,
            upload_files=[("filetoupload", xml)],
            extra_environ=dict(
                FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin
            ),
        )
        assert "FS_error" not in res.headers
        # A submission in testing has no row in submission, so the document
        # is found as the newest of the form.
        connection.indices.refresh(index=index_name)
        hit = connection.search(
            index=index_name,
            query={
                "bool": {
                    "filter": [
                        {"term": {"project_id": test_object.projectID}},
                        {"term": {"form_id": FORM}},
                    ]
                }
            },
            sort=[{"_submitted_date": "desc"}],
            size=1,
        )["hits"]["hits"][0]
        return hit["_source"], hit["_id"]

    # --- The form records its shapes at upload --------------------------------
    xlsx = os.path.join(work, FORM + ".xlsx")
    _workbook(xlsx)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "i_d"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    assert _query(
        config,
        "SELECT form_geopoint, form_geoshape FROM odkform WHERE project_id = '{}' "
        "AND form_id = '{}'".format(test_object.projectID, FORM),
    )[0] == ("home", "plot:geoshape,route:geotrace")
    testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(login, project, FORM),
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )

    # --- In testing: a bow-tie, repaired, counted by a point inside it ---------
    document, bowtie_id = push("G1", plot=BOWTIE)
    assert document["_geotype"] == "shape"
    assert document["_geoshape"]["type"] == "MultiPolygon"
    assert len(document["_geoshape"]["coordinates"]) == 2
    assert _inside(document), document

    # --- In the repository ---------------------------------------------------
    testapp.post(
        "/user/{}/project/{}/form/{}/repository/create".format(login, project, FORM),
        {"form_pkey": "i_d", "start_stage1": ""},
        status=302,
    )
    schema = None
    for _ in range(100):
        schema = get_form_details(config, test_object.projectID, FORM)["form_schema"]
        if (
            schema
            and _query(
                config,
                "SELECT COUNT(*) FROM information_schema.TABLES WHERE TABLE_SCHEMA = '{}' "
                "AND TABLE_NAME = 'maintable'".format(schema),
            )[0][0]
        ):
            break
        time.sleep(3)
    assert schema, "the geo_index repository did not build"

    # A geopoint still places the document; the shape is drawn.
    document, _ = push("G2", home=HOME, plot=SQUARE)
    assert document["_geotype"] == "shape"
    assert document["_geoshape"]["type"] == "Polygon"
    assert (
        float(document["_geolocation"]["lat"]),
        float(document["_geolocation"]["lon"]),
    ) == (-2.005, 37.005)

    # A trace alone: a line, counted by a point on it.
    document, route_id = push("G3", route=ROUTE)
    assert document["_geotype"] == "trace"
    assert document["_geoshape"]["type"] == "LineString"
    assert document["_geoshape"]["coordinates"][0] == [38.0, -3.0]
    assert _inside(document), document

    # Nothing to place: no location at all.
    document, _ = push("G4")
    assert "_geolocation" not in document and "_geotype" not in document

    # --- What the map will ask: the shapes that cross a box --------------------
    connection.indices.refresh(index=index_name)
    hits = connection.search(
        index=index_name,
        query={
            "bool": {
                "filter": [
                    {"term": {"project_id": test_object.projectID}},
                    {"term": {"form_id": FORM}},
                    {
                        "geo_shape": {
                            "_geoshape": {
                                "shape": {
                                    "type": "envelope",
                                    "coordinates": [[35.9, -0.9], [36.1, -1.1]],
                                },
                                "relation": "intersects",
                            }
                        }
                    },
                ]
            }
        },
    )["hits"]["hits"]
    assert [hit["_id"] for hit in hits] == [bowtie_id], hits

    # --- index_geometries gives an old document its shape ----------------------
    connection.update(
        index=index_name,
        id=route_id,
        script={
            "source": "ctx._source.remove('_geoshape'); ctx._source.remove('_geotype')"
        },
    )
    assert "_geoshape" not in connection.get(index=index_name, id=route_id)["_source"]
    assert (
        index_geometries(
            [
                "--project",
                test_object.projectID,
                "--form",
                FORM,
                config["global:config:file"],
            ]
        )
        == 0
    )
    document = connection.get(index=index_name, id=route_id)["_source"]
    assert document["_geotype"] == "trace"
    assert document["_geoshape"]["type"] == "LineString"
    connection.close()

    # --- A form located by a trace alone -------------------------------------
    # The pages called a form geo-referenced only when it had a geopoint, so
    # one with a trace or a shape alone had no map at all.
    xlsx = os.path.join(work, ROUTE_FORM + ".xlsx")
    _route_workbook(xlsx)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "i_d"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(login, project, ROUTE_FORM),
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )
    xml = os.path.join(work, "R1.xml")
    with open(xml, "w", encoding="utf-8") as a_file:
        a_file.write(
            "<?xml version='1.0'?>"
            '<data id="{}" version="1"><i_d>R1</i_d><route>{}</route>'
            "<meta><instanceID>uuid:{}</instanceID></meta></data>".format(
                ROUTE_FORM, ROUTE, uuid.uuid4()
            )
        )
    res = testapp.post(
        "/user/{}/project/{}/push".format(login, project),
        status=201,
        upload_files=[("filetoupload", xml)],
        extra_environ=dict(
            FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin
        ),
    )
    assert "FS_error" not in res.headers
    create_connection(config).indices.refresh(index=index_name)
    res = testapp.get(
        "/user/{}/project/{}/form/{}".format(login, project, ROUTE_FORM), status=200
    )
    assert b'id="projmap"' in res.body
    assert b"This form is not geo-referenced" not in res.body
    assert b"Geo-referenced" in res.body
    points = testapp.get(
        "/user/{}/project/{}/form/{}/get/gpspoints".format(login, project, ROUTE_FORM),
        status=200,
    ).json["points"]
    assert len(points) == 1, points
    on_route = Point(float(points[0]["long"]), float(points[0]["lati"]))
    route_line = shape(
        {
            "type": "LineString",
            "coordinates": [[38, -3], [38.01, -3.01], [38.03, -3.02]],
        }
    )
    assert route_line.distance(on_route) < 1e-9, points
