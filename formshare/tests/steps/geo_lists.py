"""GeoJSON published lists (feature 6): each case a feature on a map.

A registration form records each farm's homestead (a geopoint) and its
boundary (a geoshape); a visit form selects the farm on a map, from
plots.geojson. Through the web routes:

- the wizard offers the geopoints, geotraces and geoshapes of the table, and
  refuses a GeoJSON list that is keyed, that has no location, or that is
  named for the other format; it publishes one drawn from the boundary;
- no served column takes the name the location travels under;
- the sample is caselistgen's FeatureCollection: one feature per farm with a
  boundary, its id the farm's rowuuid, its label again as title, longitude
  first, the ring closed; the farm without one is left out;
- the visit form attaches the sample as its placeholder, and its select
  from the list is the case link: the repository hangs each visit off a
  farm. Its build empties the list's copy, as a CSV list's is cut to its
  header. A jxformtomysql from before rstools.md 27 (2026-10-07) refuses
  that with 28, and CI's image may have one, so the step asks the installed
  tool first; with an older one it checks that the page says the server
  needs a newer RSTools;
- the manifest serves the generated list in place of the placeholder, the
  farm left out is counted, and lists.xml says where the location is and
  what kind it is;
- a visit selecting a farm by its feature id is stored against that farm
  (once the form can be built);
- the list moves to the homestead, and the next pull serves points.
"""

import json
import os
import re
import shutil
import subprocess
import time
import uuid
from urllib.parse import urlparse

import openpyxl
from lxml import etree
from sqlalchemy import create_engine
from sqlalchemy.pool import NullPool

from formshare.processes.odk.api import EMPTY_FEATURE_COLLECTION
from .sql import get_form_details

FARMS = "geo_farms"
VISITS = "geo_visits"
SQUARE = "-2 37 0 0;-2 37.01 0 0;-2.01 37.01 0 0;-2.01 37 0 0;-2 37 0 0"
TRIANGLE = "-3 38 0 0;-3 38.01 0 0;-3.01 38 0 0;-3 38 0 0"
OPENROSA_MANIFEST = "{http://openrosa.org/xforms/xformsManifest}"


def _farms_workbook(path):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "farm_id", "Farm ID"])
    survey.append(["text", "farmer", "Farmer"])
    survey.append(["geopoint", "home", "Homestead"])
    survey.append(["geoshape", "boundary", "Boundary"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Farms on a map", FARMS, "1"])
    wb.save(path)


def _visits_workbook(path):
    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "visit_id", "Visit ID"])
    survey.append(["select_one_from_file plots.geojson", "farm", "Farm"])
    survey.append(["text", "notes", "Notes"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Visits to a farm on a map", VISITS, "1"])
    wb.save(path)


def _query(config, sql):
    engine = create_engine(config["sqlalchemy.url"], poolclass=NullPool)
    try:
        return engine.execute(sql).fetchall()
    finally:
        engine.dispose()


def _wait_for_build(test_object, form_id):
    config = test_object.server_config
    for _ in range(100):
        schema = get_form_details(config, test_object.projectID, form_id)["form_schema"]
        if (
            schema
            and _query(
                config,
                "SELECT COUNT(*) FROM information_schema.TABLES WHERE TABLE_SCHEMA = '{}' "
                "AND TABLE_NAME = 'maintable'".format(schema),
            )[0][0]
        ):
            return schema
        time.sleep(3)
    raise AssertionError("the repository of {} did not build".format(form_id))


def _as_device(test_object):
    return dict(FS_for_testing="true", FS_user_for_testing=test_object.assistantLogin)


def _manifest(test_object, form_id):
    return test_object.testapp.get(
        "/user/{}/project/{}/{}/manifest".format(
            test_object.randonLogin, test_object.project, form_id
        ),
        status=200,
        extra_environ=_as_device(test_object),
    ).body


def _served(test_object, manifest, file_name):
    """A manifest entry downloaded through its own downloadUrl, as a device."""
    root = etree.fromstring(manifest)
    for entry in root.findall(OPENROSA_MANIFEST + "mediaFile"):
        if entry.findtext(OPENROSA_MANIFEST + "filename") == file_name:
            return test_object.testapp.get(
                urlparse(entry.findtext(OPENROSA_MANIFEST + "downloadUrl")).path,
                status=200,
                extra_environ=_as_device(test_object),
            ).body
    raise AssertionError("{} is not in the manifest".format(file_name))


def _push(test_object, work, xml):
    path = os.path.join(work, "{}.xml".format(uuid.uuid4().hex[:8]))
    with open(path, "w", encoding="utf-8") as a_file:
        a_file.write(xml)
    res = test_object.testapp.post(
        "/user/{}/project/{}/push".format(test_object.randonLogin, test_object.project),
        status=201,
        upload_files=[("filetoupload", path)],
        extra_environ=_as_device(test_object),
    )
    assert "FS_error" not in res.headers


def _farm(farm_id, farmer, home, boundary=""):
    return (
        "<?xml version='1.0'?>"
        '<data id="{form}" version="1">'
        "<farm_id>{farm_id}</farm_id><farmer>{farmer}</farmer>"
        "<home>{home}</home><boundary>{boundary}</boundary>"
        "<meta><instanceID>uuid:{uuid}</instanceID></meta></data>"
    ).format(
        form=FARMS,
        farm_id=farm_id,
        farmer=farmer,
        home=home,
        boundary=boundary,
        uuid=uuid.uuid4(),
    )


def _assign_assistant(test_object, form_id):
    test_object.testapp.post(
        "/user/{}/project/{}/form/{}/assistants/add".format(
            test_object.randonLogin, test_object.project, form_id
        ),
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
            "coll_can_clean": "1",
        },
        status=302,
    )


def tool_reads_an_empty_geojson(odktools_path, work):
    """True when the installed jxformtomysql reads a FeatureCollection with
    no features as a CSV with only its header -- no values, no lookup, the
    select external (rstools.md 27, 2026-10-07) -- and False when it predates
    that and refuses it with 28, in a build and in a check alike."""
    probe = os.path.join(work, "empty_geojson_probe")
    if os.path.exists(probe):
        shutil.rmtree(probe)
    os.makedirs(os.path.join(probe, "tmp"))
    try:
        survey = {
            "type": "survey",
            "name": "data",
            "title": "Probe",
            "id_string": "probe",
            "sms_keyword": "probe",
            "default_language": "default",
            "version": "1",
            "children": [
                {"type": "text", "name": "probe_id", "label": "ID"},
                {
                    "type": "select one",
                    "name": "farm",
                    "label": "Farm",
                    "itemset": "plots.geojson",
                    "parameters": {},
                },
            ],
        }
        with open(os.path.join(probe, "probe.srv"), "w") as a_file:
            json.dump(survey, a_file)
        with open(os.path.join(probe, "plots.geojson"), "w") as a_file:
            a_file.write(EMPTY_FEATURE_COLLECTION)
        run = subprocess.run(
            [
                os.path.join(odktools_path, "JXFormToMysql", "jxformtomysql"),
                "-j",
                os.path.join(probe, "probe.srv"),
                "-t",
                "maintable",
                "-v",
                "probe_id",
                "-e",
                os.path.join(probe, "tmp"),
                "-o",
                "m",
                os.path.join(probe, "plots.geojson"),
            ],
            cwd=probe,
            capture_output=True,
        )
        return run.returncode == 0
    finally:
        shutil.rmtree(probe, ignore_errors=True)


def _skipped(test_object):
    return _query(
        test_object.server_config,
        "SELECT list_skipped FROM publishedlist WHERE project_id = '{}' "
        "AND list_id = 'plots'".format(test_object.projectID),
    )[0][0]


def t_e_s_t_geo_lists(test_object):
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    config = test_object.server_config
    work = os.path.join(test_object.working_dir, "geo_lists")
    if os.path.exists(work):
        shutil.rmtree(work)
    os.makedirs(work)
    lists = "/user/{}/project/{}/caselists".format(login, project)

    # --- The farms, with a homestead and a boundary ------------------------------
    xlsx = os.path.join(work, FARMS + ".xlsx")
    _farms_workbook(xlsx)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "farm_id"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    _assign_assistant(test_object, FARMS)
    testapp.post(
        "/user/{}/project/{}/form/{}/repository/create".format(login, project, FARMS),
        {"form_pkey": "farm_id", "start_stage1": ""},
        status=302,
    )
    farms_schema = _wait_for_build(test_object, FARMS)
    _push(test_object, work, _farm("F1", "Wanjiru", "-2.005 37.005 0 0", SQUARE))
    _push(test_object, work, _farm("F2", "Otieno", "-3.002 38.002 0 0", TRIANGLE))
    # Registered without a boundary: still a farm, but nothing to draw.
    _push(test_object, work, _farm("F3", "Kiprop", "-1.5 36.5 0 0"))
    rowuuids = dict(
        _query(
            config,
            "SELECT farm_id, rowuuid FROM `{}`.maintable".format(farms_schema),
        )
    )
    assert sorted(rowuuids) == ["F1", "F2", "F3"], rowuuids

    # --- The wizard ---------------------------------------------------------------
    res = testapp.get(lists + "/fieldsof/{}/maintable".format(FARMS), status=200)
    assert res.json["geo"]["home"] == "point", res.json
    assert res.json["geo"]["boundary"] == "shape", res.json
    # The geometry RSTools derives from the boundary is not a location.
    assert "boundary_geom" not in res.json["geo"], res.json
    assert "farmer" not in res.json["geo"], res.json
    plots = {
        "list_id": "plots",
        "list_filename": "plots.geojson",
        "list_format": "geojson",
        "geometry_column": "boundary",
        "source_form": FARMS,
        "source_table": "maintable",
        "label_column": "farmer",
        "list_active": "1",
        "list_key_column": "",
    }
    for changes, refusal in (
        ({"list_filename": "plots.csv"}, b"or in .geojson for a GeoJSON list"),
        ({"list_key_column": "farmer"}, b"cannot be a list of values"),
        ({"geometry_column": "farmer"}, b"Select the geopoint, geotrace or geoshape"),
        ({"geometry_column": ""}, b"Select the geopoint, geotrace or geoshape"),
    ):
        res = testapp.post(lists + "/add", dict(plots, **changes), status=200)
        assert refusal in res.body, refusal
    res = testapp.post(lists + "/add", plots, status=302)
    assert "FS_error" not in res.headers
    assert _query(
        config,
        "SELECT list_format, geometry_column, list_key_column FROM publishedlist "
        "WHERE project_id = '{}' AND list_id = 'plots'".format(test_object.projectID),
    )[0] == ("geojson", "boundary", None)
    res = testapp.get(lists, status=200)
    assert b"GeoJSON, drawn from" in res.body and b"Sample GeoJSON" in res.body

    # The farm's ID goes out as a property; the location's name is taken.
    res = testapp.post(
        lists + "/plots/edit",
        {"add_column": "1", "column_name": "farm_id", "column_as": "farm_code"},
        status=302,
    )
    assert "FS_error" not in res.headers
    for column_as in ("fs_location", "farm_code", "label"):
        res = testapp.post(
            lists + "/plots/edit",
            {"add_column": "1", "column_name": "farmer", "column_as": column_as},
            status=200,
        )
        assert b"already serves a column called" in res.body, column_as
    # A GeoJSON list cannot become a list of values.
    res = testapp.post(
        lists + "/plots/edit",
        {"change_key": "1", "list_key_column": "farmer"},
        status=200,
    )
    assert b"cannot be a list of values" in res.body
    res = testapp.get(lists + "/plots/edit", status=200)
    assert b'name="change_geometry"' in res.body
    assert b'value="boundary" selected' in res.body

    # --- The sample: caselistgen's FeatureCollection --------------------------------
    res = testapp.get(lists + "/plots/sample", status=200)
    assert res.content_type == "application/geo+json"
    assert 'filename="plots.geojson"' in res.headers["Content-Disposition"]
    sample = json.loads(res.body)
    assert sample["type"] == "FeatureCollection"
    features = {feature["id"]: feature for feature in sample["features"]}
    assert sorted(features) == sorted([rowuuids["F1"], rowuuids["F2"]]), features
    square = features[rowuuids["F1"]]
    assert square["properties"] == {
        "label": "Wanjiru",
        "title": "Wanjiru",
        "farm_code": "F1",
    }, square
    assert square["geometry"]["type"] == "Polygon"
    ring = square["geometry"]["coordinates"][0]
    # Longitude first, and the ring comes back to where it started.
    assert ring[0] == [37, -2] and ring[0] == ring[-1], ring
    assert features[rowuuids["F2"]]["properties"]["title"] == "Otieno"
    sample_file = os.path.join(work, "plots.geojson")
    with open(sample_file, "wb") as a_file:
        a_file.write(res.body)

    # --- The visit form selects the farm on a map --------------------------------------
    xlsx = os.path.join(work, VISITS + ".xlsx")
    _visits_workbook(xlsx)
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "visit_id"},
        status=302,
        upload_files=[("xlsx", xlsx)],
    )
    assert "FS_error" not in res.headers
    # The sample is the placeholder; the registry replaces it at manifest time.
    res = testapp.post(
        "/user/{}/project/{}/form/{}/upload".format(login, project, VISITS),
        status=302,
        upload_files=[("filetoupload", sample_file)],
    )
    assert "FS_error" not in res.headers
    _assign_assistant(test_object, VISITS)
    res = testapp.get(
        "/user/{}/project/{}/form/{}/caselinks".format(login, project, VISITS),
        status=200,
    )
    assert b"plots" in res.body
    res = testapp.post(
        "/user/{}/project/{}/form/{}/caselinks".format(login, project, VISITS),
        {"case_link": "1", "list_id": "plots"},
        status=302,
    )
    assert "FS_error" not in res.headers
    visits_schema = None
    if tool_reads_an_empty_geojson(config["odktools.path"], work):
        testapp.post(
            "/user/{}/project/{}/form/{}/repository/create".format(
                login, project, VISITS
            ),
            {"form_pkey": "visit_id", "start_stage1": ""},
            status=302,
        )
        visits_schema = _wait_for_build(test_object, VISITS)
        # Each visit hangs off a farm: the case link is a foreign key to its
        # rowuuid, and the sample the form was given is no lookup table.
        assert (
            "farm",
            farms_schema,
            "maintable",
            "rowuuid",
        ) in _query(
            config,
            "SELECT COLUMN_NAME, REFERENCED_TABLE_SCHEMA, REFERENCED_TABLE_NAME, "
            "REFERENCED_COLUMN_NAME FROM information_schema.KEY_COLUMN_USAGE "
            "WHERE TABLE_SCHEMA = '{}' AND TABLE_NAME = 'maintable' "
            "AND REFERENCED_TABLE_NAME IS NOT NULL".format(visits_schema),
        )
        assert not _query(
            config,
            "SELECT TABLE_NAME FROM information_schema.TABLES WHERE "
            "TABLE_SCHEMA = '{}' AND TABLE_NAME = 'lkpfarm'".format(visits_schema),
        )
    else:
        # The copy emptied for the build is refused (28), and the page says
        # it is the server's RSTools rather than the file the owner attached.
        res = testapp.post(
            "/user/{}/project/{}/form/{}/repository/create".format(
                login, project, VISITS
            ),
            {"form_pkey": "visit_id", "start_stage1": ""},
            status=200,
        )
        assert b"selects from the published GeoJSON list plots.geojson" in res.body
        assert b"does not have features" not in res.body
        assert not get_form_details(config, test_object.projectID, VISITS)[
            "form_schema"
        ]

    # --- The manifest serves the list, generated ---------------------------------------
    manifest = _manifest(test_object, VISITS)
    served = json.loads(_served(test_object, manifest, "plots.geojson"))
    assert sorted(feature["id"] for feature in served["features"]) == sorted(
        [rowuuids["F1"], rowuuids["F2"]]
    )
    # The farm without a boundary was left out, and the list says so.
    assert _skipped(test_object) == 1
    res = testapp.get(lists, status=200)
    assert b"1 without a location" in res.body
    res = testapp.get(lists + "/plots/edit", status=200)
    assert b'id="list-skipped"' in res.body
    the_list = etree.fromstring(_served(test_object, manifest, "lists.xml")).find(
        "list[@id='plots']"
    )
    assert the_list is not None
    assert (
        the_list.get("file"),
        the_list.get("kind"),
        the_list.get("format"),
        the_list.get("location"),
        the_list.get("geometry"),
        the_list.get("selector"),
        the_list.get("link"),
    ) == ("plots.geojson", "row", "geojson", "fs_location", "shape", "farm", "true")
    assert re.search(r"`boundary` AS `fs_location`", the_list.findtext("select"))

    # --- A visit to a farm chosen on the map ------------------------------------------
    _push(
        test_object,
        work,
        (
            "<?xml version='1.0'?>"
            '<data id="{}" version="1"><visit_id>V1</visit_id>'
            "<farm>{}</farm><notes>Fenced</notes>"
            "<meta><instanceID>uuid:{}</instanceID></meta></data>"
        ).format(VISITS, rowuuids["F1"], uuid.uuid4()),
    )
    if visits_schema:
        assert _query(
            config, "SELECT visit_id, farm FROM `{}`.maintable".format(visits_schema)
        ) == [("V1", rowuuids["F1"])]

    # --- Drawn from the homestead instead ----------------------------------------------
    res = testapp.post(
        lists + "/plots/edit",
        {"change_geometry": "1", "geometry_column": "farmer"},
        status=200,
    )
    assert b"Select the geopoint, geotrace or geoshape" in res.body
    res = testapp.post(
        lists + "/plots/edit",
        {"change_geometry": "1", "geometry_column": "home"},
        status=302,
    )
    assert "FS_error" not in res.headers
    manifest = _manifest(test_object, VISITS)
    served = json.loads(_served(test_object, manifest, "plots.geojson"))
    points = {feature["id"]: feature["geometry"] for feature in served["features"]}
    # Every farm has a homestead: none is left out now.
    assert sorted(points) == sorted(rowuuids.values())
    assert points[rowuuids["F1"]] == {"type": "Point", "coordinates": [37.005, -2.005]}
    assert _skipped(test_object) == 0
    the_list = etree.fromstring(_served(test_object, manifest, "lists.xml")).find(
        "list[@id='plots']"
    )
    assert the_list.get("geometry") == "point"

    # --- What the export says -------------------------------------------------------------
    res = testapp.get(lists + "/plots/export?format=json", status=200)
    exported = json.loads(res.body)["lists"][0]
    assert (exported["format"], exported["geometry_column"]) == ("geojson", "home")
