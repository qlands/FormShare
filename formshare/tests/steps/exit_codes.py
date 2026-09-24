"""Forms the form checker refuses with codes no other form in the suite
produces.

The checker is jxformtomysql. Each exit code has a handler in
formshare/processes/odk/api.py that turns the checker's XML report into the
message the user reads, and the forms here are the smallest that make the
checker return that code. The checker still returns them: pyxform rewrites
"or other" selects before the checker sees them, so code 35 is not among
these, and codes 2, 21 and 29 are no longer returned at all. Code 37 (an
option value the repository cannot keep as written) arrived on 2026-09-19.

Code 38 (a file the tool could not write, rstools.md 17.2) is no fault of a
form: the step makes a file of the repository unwritable instead, and only
when the installed tool knows the code.
"""

import os

import openpyxl

from .sql import get_form_details
from .unwritable_output import (
    make_unwritable,
    make_writable,
    tool_reports_unwritable_files,
)


def _form_with_a_bad_option_value(working_dir):
    """The smallest XLSForm whose choices carry a value the repository cannot
    keep as written: a semicolon. Written at run time so the fault is in
    plain sight rather than inside a binary file."""
    book = openpyxl.Workbook()
    survey = book.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "hid", "Household id"])
    survey.append(["select_one zones", "zone", "Zone"])
    choices = book.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    choices.append(["zones", "north", "North"])
    choices.append(["zones", "a;b", "A or B"])
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Bad option value", "bad_option_value", "1"])
    path = os.path.join(working_dir, "bad_option_value.xlsx")
    book.save(path)
    return path


def _form_to_build(working_dir):
    """A form with nothing wrong with it, to build a repository for."""
    book = openpyxl.Workbook()
    survey = book.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    survey.append(["text", "hid", "Household id"])
    survey.append(["integer", "members", "Members"])
    settings = book.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Unwritable repository", "unwritable_repository", "1"])
    path = os.path.join(working_dir, "unwritable_repository.xlsx")
    book.save(path)
    return path


def _unwritable_repository_file(test_object):
    """38 on the repository page: the page names the file and the reason,
    never the server's path, and builds nothing."""
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    form_url = "/user/{}/project/{}/form/{}".format(
        login, project, "unwritable_repository"
    )
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "hid"},
        status=302,
        upload_files=[("xlsx", _form_to_build(test_object.working_dir))],
    )
    assert "FS_error" not in res.headers
    res = testapp.post(
        form_url + "/assistants/add",
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    directory = get_form_details(
        test_object.server_config, test_object.projectID, "unwritable_repository"
    )["form_directory"]
    manifest = os.path.join(
        test_object.server_config["repository.path"],
        "odk",
        "forms",
        directory,
        "repository",
        "manifest.xml",
    )
    make_unwritable(manifest)
    try:
        res = testapp.post(
            form_url + "/repository/create",
            {"form_pkey": "hid", "start_stage1": ""},
            status=200,
        )
    finally:
        make_writable(manifest)
    assert b"FormShare could not write the files of the repository" in res.body
    assert b"not with your form" in res.body
    assert b"manifest.xml" in res.body and b"Is a directory" in res.body
    assert manifest.encode("utf-8") not in res.body
    details = get_form_details(
        test_object.server_config, test_object.projectID, "unwritable_repository"
    )
    assert details["form_schema"] is None and details["form_reptask"] is None

    res = testapp.post(form_url + "/delete", status=302)
    assert "FS_error" not in res.headers


def t_e_s_t_exit_codes(test_object):
    login = test_object.randonLogin
    project = test_object.project
    testapp = test_object.testapp
    resources = os.path.join(test_object.path, "resources", "forms", "exit_codes")

    # 20: a question called rowuuid, a name FormShare keeps for itself, and
    # one with a hyphen, which no column can carry
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "hid"},
        status=302,
        upload_files=[("xlsx", os.path.join(resources, "reserved_names.xlsx"))],
    )
    assert "FS_error" in res.headers

    # 37: an option value the repository could not keep as written, refused
    # before any repository exists (rstools.md 8.9). The message names the
    # option, the list and the problem.
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "hid"},
        status=302,
        upload_files=[("xlsx", _form_with_a_bad_option_value(test_object.working_dir))],
    )
    assert "FS_error" in res.headers
    res = res.follow(status=200)
    assert b"cannot be stored as it is written" in res.body
    assert b"a;b" in res.body and b"semicolon" in res.body

    # 14: a CSV whose column heading has a space in it. The form is accepted
    # and so is the file; the check that runs once the file is there is not
    form_url = "/user/{}/project/{}/form/{}".format(login, project, "bad_csv_form")
    res = testapp.post(
        "/user/{}/project/{}/forms/add".format(login, project),
        {"form_pkey": "hid"},
        status=302,
        upload_files=[("xlsx", os.path.join(resources, "bad_csv.xlsx"))],
    )
    assert "FS_error" not in res.headers
    # The check is only shown once someone could collect data with the form
    res = testapp.post(
        form_url + "/assistants/add",
        {
            "coll_id": "{}|{}|{}".format(
                test_object.projectID,
                test_object.assistantLogin,
                test_object.assistantLoginUUID,
            ),
            "coll_can_submit": "1",
        },
        status=302,
    )
    assert "FS_error" not in res.headers
    res = testapp.get(form_url, status=200)
    assert b"Repository check pending" in res.body
    res = testapp.post(
        form_url + "/upload",
        status=302,
        upload_files=[("filetoupload", os.path.join(resources, "bad.csv"))],
    )
    assert "FS_error" not in res.headers
    res = testapp.get(form_url, status=200)
    assert b"This form cannot create a repository" in res.body
    assert b"invalid characters in column headers" in res.body

    res = testapp.post(form_url + "/delete", status=302)
    assert "FS_error" not in res.headers

    # 38: a file of the repository could not be written (rstools.md 17.2)
    if tool_reports_unwritable_files(
        test_object.server_config["odktools.path"], test_object.working_dir
    ):
        _unwritable_repository_file(test_object)
